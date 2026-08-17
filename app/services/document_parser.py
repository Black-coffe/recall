"""Парсинг документів → чистий текст (+ легка структура) для RAG (Phase 16A).

Документ перетворюється на plain-text (markdown-подібний), який кладеться у
`transcriptions.transcript_text`. Далі ВЕСЬ нижній стек працює без змін:
чанкінг (embeddings._chunk_from_text), локальні embeddings (e5-large), FTS5,
Claude-enrichment (summary/сутності/action items), граф, напрямки, RAG-чат.

Парсери реєструються по розширенню. Усе локально й безкоштовно:
  .md / .markdown / .txt — пряме читання (UTF-8, errors=replace)
  .docx                  — python-docx (вже в deps): параграфи + таблиці → markdown
  .pdf                   — PyMuPDF (fitz): постранічний витяг тексту (+ блоки сторінок)
  .pptx                  — python-pptx: текст слайдів + нотатки (+ блоки слайдів)
  .xlsx                  — openpyxl: аркуші → markdown-таблиці (+ блоки аркушів)
  .csv                   — вбудований csv: одна таблиця (без залежності)
  .png/.jpg/.tiff/...    — Tesseract OCR (pytesseract): текст із зображення
  скан-PDF               — OCR-фолбек постранічно (рендер fitz → tesseract)

Провенанс (Phase 16B/16C): PDF/PPTX/XLSX/CSV повертають `blocks` =
[{text, page, section, ...}], де page — номер сторінки/слайда/аркуша. Це дозволяє
RAG-цитатам вказувати «стор. 3 / слайд 5 / лист «Бюджет»». Формати без структури
(md/txt/docx) повертають blocks=None.

Таблиці (Phase 16C): голий дамп ячейок погано ембедиться. Тому окрім markdown
(для FTS + структурного вектора) над зразком кожного аркуша Claude генерує
NL-опис «що в таблиці» (text_polishing.describe_sheets, у documents.py) — він
домішується у текст блоку → таблиця стає знаходимою семантично. Опис опційний:
без ANTHROPIC_API_KEY лишаються тільки markdown-таблиці.

OCR (Phase 16D): зображення і скан-PDF (без текстового шару) розпізнаються
Tesseract'ом (pytesseract — обгортка над СИСТЕМНИМ tesseract.exe, БЕЗ torch/
transformers, тому не чіпає піни — на відміну від surya/docTR). Потребує
встановленого tesseract + мовних пакетів. ocr_available()=False → зображення
кидають ParserUnavailable, скан-PDF → needs_ocr, решта парсингу працює.

Lazy guarded imports: модуль імпортується навіть без PyMuPDF/python-pptx/openpyxl/
pytesseract — ці формати тоді кидають ParserUnavailable, а .md/.txt/.docx/.csv
працюють завжди (python-docx у базових deps, csv вбудований).
"""
from __future__ import annotations

import hashlib
import logging
import os
import re
from typing import Callable, Optional

from app.utils.files import ALLOWED_DOCUMENT_EXTENSIONS


logger = logging.getLogger(__name__)

# Бамп при зміні логіки парсингу → re-parse застарілих (16E).
PARSER_VERSION = 1

# Якщо у текстовому шарі PDF менше символів на сторінку — ймовірно скан без
# текстового шару (потрібен OCR, Phase 16D). Помічаємо у meta, але текст
# все одно повертаємо (раптом є частковий шар).
_PDF_MIN_CHARS_PER_PAGE = 24

# --- OCR (Phase 16D) ---
# Tesseract через pytesseract — тонка обгортка над системним бінарём, БЕЗ
# torch/transformers (на відміну від surya/docTR) → не чіпає крихку цепочку
# пінів. Потребує встановленого tesseract.exe + мовні пакети у PATH.
_OCR_LANG = os.environ.get("OCR_LANG", "ukr+rus+eng")
_OCR_DPI = int(os.environ.get("OCR_DPI", "200"))      # рендер сторінки PDF для OCR
_OCR_MAX_PAGES = int(os.environ.get("OCR_MAX_PAGES", "50"))  # кап синхронного OCR
_ocr_unavailable_reason: Optional[str] = None


class DocumentParseError(Exception):
    """Не вдалося розібрати документ (порожній / биткий / непідтримуваний)."""


class ParserUnavailable(DocumentParseError):
    """Парсер для формату недоступний (напр. PyMuPDF не встановлено для PDF)."""


def ext_of(filename: str) -> str:
    if not filename or '.' not in filename:
        return ''
    return filename.rsplit('.', 1)[1].lower()


def is_supported(filename: str) -> bool:
    return ext_of(filename) in ALLOWED_DOCUMENT_EXTENSIONS


def content_hash(text: str) -> str:
    """sha256 нормалізованого тексту — для дедупу повторних завантажень (16E)."""
    return hashlib.sha256((text or "").encode("utf-8", errors="replace")).hexdigest()


def _normalize_text(text: str) -> str:
    """Уніфікувати переноси рядків і схлопнути >2 порожніх рядки поспіль."""
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)        # трейлінг-пробіли
    text = re.sub(r"\n{3,}", "\n\n", text)         # ≤1 порожній рядок між блоками
    return text.strip()


# ============================================================
# OCR (Phase 16D) — graceful, як embeddings.is_available()
# ============================================================

def ocr_available() -> bool:
    """Чи доступний OCR: pytesseract імпортується І системний tesseract присутній."""
    global _ocr_unavailable_reason
    try:
        import pytesseract  # noqa: F401
        from PIL import Image  # noqa: F401
    except Exception as e:  # pragma: no cover
        _ocr_unavailable_reason = f"pytesseract/Pillow недоступні: {e}"
        return False
    try:
        import pytesseract
        # get_tesseract_version() іде повз subprocess_args() самого pytesseract
        # (там SW_HIDE) — без патчу під pythonw.exe блимає консоль.
        silence_pytesseract_console_windows()
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        _ocr_unavailable_reason = (
            "системний Tesseract не знайдено. Встановіть Tesseract OCR "
            "(+ мовні пакети ukr/rus/eng) і додайте його у PATH."
        )
        return False


def ocr_unavailability_reason() -> str:
    return _ocr_unavailable_reason or "OCR недоступний"


def _ocr_image_obj(img) -> str:
    """OCR одного PIL-зображення. Пробує _OCR_LANG, з фолбеком на дефолтну мову
    (якщо бракує мовного пакету tesseract кидає TesseractError)."""
    import pytesseract
    try:
        return pytesseract.image_to_string(img, lang=_OCR_LANG)
    except pytesseract.TesseractError:
        try:
            return pytesseract.image_to_string(img)  # дефолтна мова (eng)
        except Exception:
            return ""
    except Exception:
        return ""


def _ocr_pdf_page(page) -> str:
    """Відрендерити сторінку PDF у растр і прогнати через OCR."""
    try:
        import io
        from PIL import Image
        pix = page.get_pixmap(dpi=_OCR_DPI)
        with Image.open(io.BytesIO(pix.tobytes("png"))) as img:
            return _normalize_text(_ocr_image_obj(img))
    except Exception as e:  # pragma: no cover
        logger.warning("[ocr] сторінку не розпізнано: %s", e)
        return ""


# ============================================================
# Парсери по форматах
# ============================================================

def _parse_text(path: str) -> dict:
    """.md / .markdown / .txt — пряме читання."""
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        raw = f.read()
    return {"text": _normalize_text(raw), "page_count": None, "meta": {}}


def _docx_iter_blocks(doc):
    """Ітерувати параграфи і таблиці у порядку появи в документі."""
    from docx.document import Document as _DocxDocument
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table as _Table
    from docx.text.paragraph import Paragraph as _Paragraph

    body = doc.element.body
    for child in body.iterchildren():
        if isinstance(child, CT_P):
            yield ("p", _Paragraph(child, doc))
        elif isinstance(child, CT_Tbl):
            yield ("table", _Table(child, doc))


def _docx_heading_level(paragraph) -> Optional[int]:
    """Рівень заголовка (1..6) якщо стиль 'Heading N' / 'Заголовок N', інакше None."""
    try:
        name = (paragraph.style.name or "")
    except Exception:
        return None
    m = re.search(r"(\d+)$", name)
    if m and ("Heading" in name or "Заголовок" in name or "Title" in name):
        return min(int(m.group(1)), 6)
    if name in ("Title", "Заголовок"):
        return 1
    return None


def _docx_table_to_markdown(table) -> str:
    rows = []
    for row in table.rows:
        cells = [" ".join((c.text or "").split()) for c in row.cells]
        rows.append("| " + " | ".join(cells) + " |")
    if not rows:
        return ""
    ncol = len(table.rows[0].cells)
    sep = "| " + " | ".join(["---"] * ncol) + " |"
    return "\n".join([rows[0], sep, *rows[1:]])


def _parse_docx(path: str) -> dict:
    """.docx — параграфи (із заголовками) + таблиці → markdown."""
    try:
        import docx  # python-docx
    except ImportError as e:  # pragma: no cover
        raise ParserUnavailable(f"python-docx недоступний: {e}") from e

    doc = docx.Document(path)
    parts: list[str] = []
    for kind, block in _docx_iter_blocks(doc):
        if kind == "p":
            txt = (block.text or "").strip()
            if not txt:
                continue
            level = _docx_heading_level(block)
            parts.append(("#" * level + " " + txt) if level else txt)
        elif kind == "table":
            md = _docx_table_to_markdown(block)
            if md:
                parts.append(md)
    text = _normalize_text("\n\n".join(parts))
    return {"text": text, "page_count": None, "meta": {}}


def _parse_pdf(path: str) -> dict:
    """.pdf — постранічний витяг тексту через PyMuPDF (fitz)."""
    try:
        import fitz  # PyMuPDF
    except ImportError as e:
        raise ParserUnavailable(
            "PyMuPDF не встановлено — PDF-парсинг недоступний. "
            "Встановіть: pip install PyMuPDF"
        ) from e

    ocr_on = ocr_available()  # Phase 16D: фолбек для сканів без текстового шару
    pages: list[str] = []
    blocks: list[dict] = []
    ocr_pages = 0
    with fitz.open(path) as doc:
        page_count = doc.page_count
        for i, page in enumerate(doc):
            ptext = _normalize_text(page.get_text("text") or "")
            # Сторінка майже без тексту + OCR доступний → розпізнаємо растром.
            if len(ptext) < _PDF_MIN_CHARS_PER_PAGE and ocr_on and ocr_pages < _OCR_MAX_PAGES:
                otext = _ocr_pdf_page(page)
                if len(otext) > len(ptext):
                    ptext = otext
                    ocr_pages += 1
            pages.append(ptext)
            if ptext:
                # Провенанс (Phase 16B): кожна непорожня сторінка — окремий блок.
                blocks.append({"text": ptext, "page": i + 1, "section": None})
    text = _normalize_text("\n\n".join(p for p in pages if p))

    meta = {}
    if ocr_pages:
        meta["ocr_pages"] = ocr_pages
    # Скан без текстового шару І OCR недоступний → позначаємо (повідомлення в parse_document).
    if page_count and len(text) < _PDF_MIN_CHARS_PER_PAGE * page_count and not ocr_on:
        meta["needs_ocr"] = True
    return {"text": text, "page_count": page_count, "blocks": blocks, "meta": meta}


def _parse_pptx(path: str) -> dict:
    """.pptx — текст слайдів + нотатки доповідача (Phase 16B). page=номер слайда,
    section=заголовок слайда (для провенансу цитат)."""
    try:
        from pptx import Presentation  # python-pptx
    except ImportError as e:
        raise ParserUnavailable(
            "python-pptx не встановлено — PPTX-парсинг недоступний. "
            "Встановіть: pip install python-pptx"
        ) from e

    prs = Presentation(path)
    slides = list(prs.slides)
    blocks: list[dict] = []
    parts: list[str] = []
    for i, slide in enumerate(slides):
        # Заголовок шукаємо за shape_id (НЕ за `is`: python-pptx віддає нові
        # обгортки на кожній ітерації, тож identity-порівняння не спрацьовує).
        title_id = None
        try:
            _ts = slide.shapes.title
            title_id = _ts.shape_id if _ts is not None else None
        except Exception:
            title_id = None
        title: Optional[str] = None
        lines: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            txt = (shape.text_frame.text or "").strip()
            if not txt:
                continue
            if title is None and title_id is not None:
                try:
                    if shape.shape_id == title_id:
                        title = txt.splitlines()[0][:120]
                except Exception:
                    pass
            lines.append(txt)
        # нотатки доповідача
        notes = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        except Exception:
            notes = ""
        slide_text = "\n".join(lines)
        if notes:
            slide_text += f"\n\n[Нотатки доповідача]\n{notes}"
        slide_text = _normalize_text(slide_text)
        if slide_text:
            blocks.append({"text": slide_text, "page": i + 1, "section": title})
            parts.append(slide_text)

    text = _normalize_text("\n\n".join(parts))
    return {"text": text, "page_count": len(slides), "blocks": blocks, "meta": {}}


# ------------------------------------------------------------
# Таблиці: XLSX / CSV (Phase 16C)
# ------------------------------------------------------------
# Кап рядків на аркуш для збереженого markdown — щоб гігантський дамп не роздув
# transcript_text. Опис («що в таблиці») робить Claude (text_polishing.describe_sheets)
# над зразком — тому повний обсяг для семантики не потрібен.
_TABLE_MAX_ROWS = 1000


def _md_cell(v) -> str:
    """Значення ячейки → безпечний markdown-текст (екран pipe, без переносів)."""
    if v is None:
        return ""
    s = str(v).replace("\r", " ").replace("\n", " ").replace("|", "\\|")
    return " ".join(s.split())


def _rows_to_markdown(rows: list[list]) -> tuple[str, int, int]:
    """rows (перший = заголовок) → markdown-таблиця. Returns (md, n_rows, n_cols)."""
    rows = [r for r in rows if r is not None]
    if not rows:
        return "", 0, 0
    ncols = max((len(r) for r in rows), default=0)
    if ncols == 0:
        return "", 0, 0
    norm = [[_md_cell(r[i]) if i < len(r) else "" for i in range(ncols)] for r in rows]
    header = norm[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * ncols) + " |",
    ]
    for r in norm[1:]:
        lines.append("| " + " | ".join(r) + " |")
    return "\n".join(lines), len(rows), ncols


def _parse_xlsx(path: str) -> dict:
    """.xlsx — кожен аркуш → markdown-таблиця як окремий блок (page=індекс
    аркуша, section=назва аркуша). data_only=True → значення формул, не формули."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise ParserUnavailable(
            "openpyxl не встановлено — XLSX-парсинг недоступний. "
            "Встановіть: pip install openpyxl"
        ) from e

    wb = load_workbook(path, read_only=True, data_only=True)
    blocks: list[dict] = []
    parts: list[str] = []
    try:
        for idx, ws in enumerate(wb.worksheets):
            rows: list[list] = []
            truncated = False
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= _TABLE_MAX_ROWS:
                    truncated = True
                    break
                if row is None or all(c is None for c in row):
                    continue
                rows.append(list(row))
            if not rows:
                continue
            md, nrows, ncols = _rows_to_markdown(rows)
            if not md:
                continue
            if truncated:
                md += f"\n\n_(показано перші {_TABLE_MAX_ROWS} рядків)_"
            title = ws.title or f"Аркуш {idx + 1}"
            sheet_text = _normalize_text(f"## Аркуш «{title}»\n\n{md}")
            blocks.append({"text": sheet_text, "page": idx + 1, "section": title,
                           "rows": nrows, "cols": ncols, "tabular": True})
            parts.append(sheet_text)
    finally:
        wb.close()

    text = _normalize_text("\n\n".join(parts))
    return {"text": text, "page_count": len(blocks), "blocks": blocks, "meta": {}}


def _parse_csv(path: str) -> dict:
    """.csv — один блок-таблиця (page=1). Делімітер визначається авто (,/;/таб)."""
    import csv as _csv

    rows: list[list] = []
    truncated = False
    with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
        sample = f.read(8192)
        f.seek(0)
        try:
            dialect = _csv.Sniffer().sniff(sample, delimiters=",;\t")
        except _csv.Error:
            dialect = _csv.excel
        reader = _csv.reader(f, dialect)
        for i, row in enumerate(reader):
            if i >= _TABLE_MAX_ROWS:
                truncated = True
                break
            rows.append(row)

    if not rows:
        return {"text": "", "page_count": 0, "blocks": [], "meta": {}}
    md, nrows, ncols = _rows_to_markdown(rows)
    if truncated:
        md += f"\n\n_(показано перші {_TABLE_MAX_ROWS} рядків)_"
    md = _normalize_text(md)
    block = {"text": md, "page": 1, "section": None,
             "rows": nrows, "cols": ncols, "tabular": True}
    return {"text": md, "page_count": 1, "blocks": [block], "meta": {}}


# ------------------------------------------------------------
# Зображення: OCR (Phase 16D)
# ------------------------------------------------------------
_IMAGE_EXTS = ("png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp")


def _parse_image(path: str) -> dict:
    """Зображення → текст через OCR (Tesseract). Без OCR — ParserUnavailable."""
    if not ocr_available():
        raise ParserUnavailable(
            "OCR недоступний — розпізнавання зображень вимкнено. "
            + ocr_unavailability_reason()
        )
    try:
        from PIL import Image
    except ImportError as e:  # pragma: no cover
        raise ParserUnavailable(f"Pillow недоступний: {e}") from e
    try:
        with Image.open(path) as img:
            text = _normalize_text(_ocr_image_obj(img))
    except Exception as e:
        raise DocumentParseError(f"Не вдалося відкрити зображення: {e}") from e
    return {"text": text, "page_count": 1, "blocks": None, "meta": {"ocr": True}}


# розширення → парсер
_PARSERS: dict[str, Callable[[str], dict]] = {
    "md": _parse_text,
    "markdown": _parse_text,
    "txt": _parse_text,
    "docx": _parse_docx,
    "pdf": _parse_pdf,
    "pptx": _parse_pptx,
    "xlsx": _parse_xlsx,
    "csv": _parse_csv,
    **{ext: _parse_image for ext in _IMAGE_EXTS},
}


# ============================================================
# Публічний API
# ============================================================

def parse_document(path: str, filename: Optional[str] = None) -> dict:
    """Розібрати документ → текст + метадані. Тип визначається за розширенням
    filename (якщо задано) інакше path.

    Returns {text, doc_type, page_count, byte_size, content_hash, char_count,
             parser_version, meta}.
    Raises DocumentParseError якщо формат не підтримується, парсер недоступний,
    або витягнутий текст порожній.
    """
    ext = ext_of(filename or path)
    parser = _PARSERS.get(ext)
    if parser is None:
        raise DocumentParseError(f"Непідтримуваний формат документа: '.{ext}'")

    try:
        byte_size = os.path.getsize(path)
    except OSError:
        byte_size = None

    result = parser(path)
    text = (result.get("text") or "").strip()
    if not text:
        meta = result.get("meta") or {}
        if meta.get("needs_ocr"):
            raise DocumentParseError(
                "У PDF не знайдено текстового шару (схоже на скан), а OCR недоступний. "
                "Встановіть Tesseract OCR (+ мовні пакети ukr/rus/eng) у PATH."
            )
        if meta.get("ocr"):
            raise DocumentParseError("OCR не знайшов тексту на зображенні.")
        raise DocumentParseError("Документ порожній або не містить тексту")

    return {
        "text": text,
        "doc_type": ext,
        "page_count": result.get("page_count"),
        "byte_size": byte_size,
        "content_hash": content_hash(text),
        "char_count": len(text),
        "parser_version": PARSER_VERSION,
        # Блоки {text, page, section} для провенансу (Phase 16B). None для
        # форматів без структури (md/txt/docx) → чанкінг падає на текстовий.
        "blocks": result.get("blocks"),
        "meta": result.get("meta") or {},
    }

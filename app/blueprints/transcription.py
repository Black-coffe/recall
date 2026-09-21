"""Transcription endpoints (Phase 5.8).

- POST   /api/transcribe                      (file/youtube → transcript)
- GET    /api/history                         (list з фільтрами + FTS)
- GET    /api/history/<id>                    (get one)
- PATCH  /api/history/<id>                    (власна назва/опис, v44)
- DELETE /api/history/<id>                    (soft-delete one, T4.6)
- POST   /api/history/<id>/restore            (undo soft-delete, T4.6)
- POST   /api/history/bulk_delete             (mass, ЗАЛИШАЄТЬСЯ фізичним)
- POST   /api/history/bulk_export             (mass JSON/TXT)
- POST   /api/transcription/<id>/polish       (Claude API)
- POST   /api/export/<format>                 (TXT/SRT/JSON одного транскрипту)
"""
from __future__ import annotations

import json
import logging
import os
import re
import time
from datetime import datetime

from flask import Blueprint, current_app, jsonify, request, send_file
from werkzeug.utils import secure_filename

from app import state
from app.repositories import transcriptions as tx_repo
from app.services import dedup_audio, record_meta, text_polishing
from app.utils.audio import extract_audio_from_video as _extract_audio_raw
from app.utils.files import allowed_file, format_srt_timestamp, is_video_file
from app.utils.fts import sanitize_fts_query
from app.utils.paths import safe_path_within, safe_path_within_any


logger = logging.getLogger(__name__)
transcription_bp = Blueprint('transcription', __name__)


def _get_db():
    from app.db.connection import get_db_connection
    return get_db_connection(current_app.config['DATABASE'])


def _make_extract_log(process_id):
    if not process_id:
        return None
    def _log(stage, message, progress=None, status="processing"):
        state.add_log(process_id, stage, message, progress, status)
    return _log


# ============================================================
# T7.1: source_type resolution — one function per /api/transcribe scenario.
#
# Each resolve_*_source() reads the current Flask request (form/files) and
# returns (filepath, source_name, meta) on success. meta is a dict with the
# fixed keys 'youtube_info', 'download_id', 'library_audio_id',
# 'library_recording_sid', 'source_type' (source_type is normally an echo of
# the input, EXCEPT resolve_library_source() which may inherit a different
# one from the audio_downloads row — the caller must use it, not its own
# local variable, from that point on).
#
# On a validation failure, they raise _SourceResolutionError carrying the
# exact (jsonify payload, status) the old inline code used to `return`
# directly — the HTTP handler below turns that back into a response. This
# keeps every original error message/status byte-for-byte while making each
# scenario callable/testable on its own (e.g. via app.test_request_context()).
# ============================================================

class _SourceResolutionError(Exception):
    """Carries a ready-to-return (jsonify payload, status) from a resolve_*_source()."""

    def __init__(self, payload: dict, status: int):
        super().__init__(payload.get('error', 'source resolution error'))
        self.payload = payload
        self.status = status


def _empty_source_meta(source_type: str) -> dict:
    return {
        'youtube_info': {},
        'download_id': None,
        'library_audio_id': None,
        'library_recording_sid': None,
        'source_type': source_type,
        # editable-title-description-02: копія з Аудіотеки (лише 'library'
        # заповнює), форма (request.form title/description) має пріоритет.
        'title': None,
        'description': None,
    }


def resolve_file_source() -> tuple[str, str, dict]:
    """source_type='file': multipart upload, extracting audio if it's a video."""
    if 'audio' not in request.files:
        raise _SourceResolutionError({"success": False, "error": "Файл не знайдено"}, 400)
    file = request.files['audio']
    if file.filename == '':
        raise _SourceResolutionError({"success": False, "error": "Файл не вибрано"}, 400)
    if not allowed_file(file.filename):
        raise _SourceResolutionError({"success": False, "error": "Непідтримуваний формат файлу"}, 400)

    filename = secure_filename(file.filename)
    if not filename:
        filename = "audio_file"
    original_ext = file.filename.rsplit('.', 1)[1].lower() if '.' in file.filename else 'mp3'
    if not filename.endswith(f'.{original_ext}'):
        filename = f"{filename}.{original_ext}"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{timestamp}_{filename}"
    upload_dir = current_app.config['UPLOAD_FOLDER']
    filepath = safe_path_within(upload_dir, os.path.join(upload_dir, filename))
    if filepath is None:
        logger.error(f"Спроба Path Traversal: {filename}")
        raise _SourceResolutionError({"success": False, "error": "Некоректне ім'я файлу"}, 400)

    file.save(filepath)
    source_name = file.filename
    logger.info(f"Файл збережено: {filepath}")

    # Якщо відео — витягаємо аудіо
    if is_video_file(file.filename):
        logger.info(f"Виявлено відео файл: {file.filename}")
        audio_filename = f"{timestamp}_extracted_audio.mp3"
        audio_filepath = os.path.join(current_app.config['UPLOAD_FOLDER'], audio_filename)
        if not _extract_audio_raw(filepath, audio_filepath, add_log=None):
            if os.path.exists(filepath):
                os.remove(filepath)
            raise _SourceResolutionError({"success": False, "error": "Не вдалося витягти аудіо з відео"}, 500)
        if os.path.exists(filepath):
            os.remove(filepath)
            logger.info(f"Відео файл видалено: {filepath}")
        filepath = audio_filepath
        logger.info(f"Використовуємо витягнуте аудіо: {filepath}")

    return filepath, source_name, _empty_source_meta('file')


def resolve_youtube_source() -> tuple[str, str, dict]:
    """source_type='youtube': previously completed /api/youtube/download by download_id."""
    download_id = request.form.get('download_id')
    if not download_id or download_id not in state.download_progress:
        raise _SourceResolutionError({"success": False, "error": "Невірний ID завантаження"}, 400)
    progress = state.download_progress.get(download_id)
    if progress['status'] != 'completed':
        raise _SourceResolutionError({"success": False, "error": "Завантаження ще не завершено"}, 400)
    filepath = progress['file_path']
    youtube_info = progress['info']
    source_name = youtube_info['title']
    state.add_log(download_id, "transcribe", "Начало транскрибации аудио...", 55, "processing")
    meta = _empty_source_meta('youtube')
    meta['youtube_info'] = youtube_info
    meta['download_id'] = download_id
    return filepath, source_name, meta


def resolve_recording_source() -> tuple[str, str, dict]:
    """source_type='recording': Phase 9.8, a finalized server-side recording session."""
    sid = request.form.get('recording_session_id')
    if not sid:
        raise _SourceResolutionError({"success": False, "error": "recording_session_id обов'язковий"}, 400)
    if state.recording_service is None:
        raise _SourceResolutionError({"success": False, "error": "Recording feature вимкнено"}, 503)
    try:
        manifest = state.recording_service.store.read(sid)
    except Exception:
        logger.debug("recording manifest read failed for sid=%s", sid, exc_info=True)
        raise _SourceResolutionError({"success": False, "error": "Recording сесія не знайдена"}, 404)
    if manifest.get('status') != 'finalized':
        raise _SourceResolutionError({"success": False, "error": "Recording ще не finalized"}, 400)
    final_mp3 = manifest.get('final_mp3_path')
    if not final_mp3 or not os.path.isfile(final_mp3):
        raise _SourceResolutionError({"success": False, "error": "Final MP3 не створений"}, 400)
    source_name = manifest.get('name') or manifest.get('auto_name') or sid
    return final_mp3, source_name, _empty_source_meta('recording')


def resolve_library_source() -> tuple[str, str, dict]:
    """source_type='library': Phase 9.10+, any Аудіотека (audio_downloads) row by id.

    Inherits the row's original source_type ('youtube'/'recording'/'file') for
    the transcriptions.source_type INSERT further downstream — the caller
    MUST use meta['source_type'] from here on, not its own local variable.
    """
    audio_id = request.form.get('audio_download_id')
    if not audio_id:
        raise _SourceResolutionError({"success": False, "error": "audio_download_id обов'язковий"}, 400)
    try:
        audio_id_int = int(audio_id)
    except (TypeError, ValueError):
        raise _SourceResolutionError({"success": False, "error": "audio_download_id мусить бути числом"}, 400)
    with _get_db() as conn:
        row = conn.execute(
            'SELECT * FROM audio_downloads WHERE id = ?', (audio_id_int,)
        ).fetchone()
    if not row:
        raise _SourceResolutionError({"success": False, "error": "Запис в Аудіотеці не знайдено"}, 404)
    filepath = row['file_path']
    if not filepath or not os.path.isfile(filepath):
        raise _SourceResolutionError({"success": False, "error": f"Файл відсутній: {filepath}"}, 400)
    source_name = row['title']
    # Успадковуємо оригінальний source_type щоб transcript у Історії
    # групувався правильно (YouTube → 'youtube', recording → 'recording')
    inherited = (row['source_type'] if 'source_type' in row.keys() else None) or 'youtube'
    meta = _empty_source_meta(inherited)
    meta['library_audio_id'] = audio_id_int
    # Успадковуємо власну назву/опис Аудіотеки — форма перебиває це в transcribe().
    meta['title'] = row['title'] if 'title' in row.keys() else None
    meta['description'] = row['description'] if 'description' in row.keys() else None
    # Phase 21: для recording зберігаємо sid → нижче лінкуємо копілот-сесію.
    if inherited == 'recording':
        meta['library_recording_sid'] = row['recording_session_id'] if 'recording_session_id' in row.keys() else None
    # Якщо youtube — заповнюємо youtube_info для INSERT
    if inherited == 'youtube':
        meta['youtube_info'] = {
            'video_id': row['youtube_id'],
            'title': row['title'],
            'author': row['author'],
            'duration': row['duration'],
            'thumbnail': row['thumbnail_url'],
            'original_url': row['youtube_url'],
        }
    return filepath, source_name, meta


_SOURCE_RESOLVERS = {
    'file': resolve_file_source,
    'youtube': resolve_youtube_source,
    'recording': resolve_recording_source,
    'library': resolve_library_source,
}


@transcription_bp.route('/api/transcribe', methods=['POST'])
def transcribe():
    """Транскрибація аудіо (file upload або YouTube via download_id)."""
    logger.info(f"=== DEBUG /api/transcribe ===")
    logger.info(f"request.form: {dict(request.form)}")
    logger.info(f"request.files keys: {list(request.files.keys())}")
    logger.info(f"content_type: {request.content_type}")
    logger.info(f"content_length: {request.content_length}")

    source_type = request.form.get('source_type', 'file')
    model_name = request.form.get('model', 'base')
    language = request.form.get('language', 'uk')

    # editable-title-description-02: власні title/description з форми —
    # валідуємо ДО початку роботи резолвера (файл/завантаження/транскрипція).
    try:
        form_title = record_meta.normalize_title(request.form.get('title'))
        form_description = record_meta.normalize_description(request.form.get('description'))
    except ValueError as e:
        return jsonify({"success": False, "error": str(e)}), 400

    resolver = _SOURCE_RESOLVERS.get(source_type)
    if resolver is None:
        return jsonify({"success": False, "error": "Невірний тип джерела"}), 400
    try:
        filepath, source_name, meta = resolver()
    except _SourceResolutionError as e:
        return jsonify(e.payload), e.status

    # Форма перебиває копію з Аудіотеки (лише 'library'-джерело її несе).
    record_title = form_title if form_title is not None else meta['title']
    record_description = form_description if form_description is not None else meta['description']

    youtube_info = meta['youtube_info']
    download_id = meta['download_id']
    library_audio_id: int | None = meta['library_audio_id']  # для active_library_transcriptions
    # Якщо транскрибуємо recording library-шляхом (audio_download_id) — sid
    # запису для привʼязки копілот-сесії беремо з рядка audio_downloads,
    # бо у формі recording_session_id немає.
    library_recording_sid: str | None = meta['library_recording_sid']
    # library-джерело може успадковувати інший source_type (T7.1) — усе, що
    # нижче (diarize/INSERT/enrichment), має орієнтуватись саме на нього.
    source_type = meta['source_type']

    if library_audio_id is not None and state.active_library_transcriptions is not None:
        _existing_job = state.active_library_transcriptions.get(library_audio_id)
        if _existing_job is not None:
            # Пост-хвильова знахідка 03.07.2026: повторний POST з тим самим
            # audio_download_id (типовий сценарій — клієнтський HTTP-таймаут
            # ретраїть запит, поки перша транскрипція ще йде синхронно всередині
            # попереднього запиту) запускав ДРУГУ повну задачу (whisper+діаризація
            # вдруге на тому самому файлі — витрачене GPU-час на годинних записах).
            # 409 замість тихого дубля.
            return jsonify({
                "success": False,
                "error": "Транскрипція цього запису вже виконується",
                "error_code": "ALREADY_RUNNING",
                "active": _existing_job,
            }), 409
        state.active_library_transcriptions[library_audio_id] = {
            'audio_download_id': library_audio_id,
            'started_at': time.time(),
            'stage': 'transcribing',
            'progress': 0,
        }

    try:
        start_time = time.time()

        def transcription_progress_callback(progress_info):
            if not isinstance(progress_info, dict):
                return
            progress_value = progress_info.get('progress', 0)
            if source_type == 'youtube' and download_id:
                message = progress_info.get('message', 'Обработка...')
                overall_progress = 55 + (progress_value * 0.4)
                segment = progress_info.get('segment')
                if segment:
                    try:
                        state.sse_broker.publish(download_id, "segment", segment)
                    except Exception:
                        logger.debug(
                            "SSE segment publish failed for %s (best-effort, не блокує transcribe)",
                            download_id, exc_info=True,
                        )
                state.add_log(download_id, "transcribe", message, overall_progress, "processing")
            if library_audio_id is not None and state.active_library_transcriptions is not None:
                entry = state.active_library_transcriptions.get(library_audio_id)
                if entry is not None:
                    entry['progress'] = progress_value

        result = state.whisper_manager.transcribe_with_progress(
            filepath, model_name, language,
            progress_callback=transcription_progress_callback,
        )
        processing_time = time.time() - start_time

        if "error" in result:
            if source_type == 'youtube' and download_id:
                state.add_log(download_id, "transcribe", f"Помилка: {result['error']}", -1, "error")
            return jsonify(result), 500

        if source_type == 'youtube' and download_id:
            state.add_log(download_id, "transcribe", "Транскрибация завершена", 95, "processing")

        result["processing_time"] = round(processing_time, 2)
        result["model_used"] = model_name
        result["source_name"] = source_name
        result["source_type"] = source_type
        if source_type == 'youtube':
            result["youtube_info"] = youtube_info

        # Phase 10.3: speaker diarization (опц., завжди для recording).
        # Помилка діарізації НЕ ламає transcribe — фічу можна вимкнути в .env
        # видаленням HF_TOKEN, тоді просто збереже сегменти без speaker-полів.
        diarize_flag = request.form.get('diarize', '').lower() in ('1', 'true', 'yes')
        should_diarize = source_type == 'recording' or diarize_flag
        if should_diarize and library_audio_id is not None and state.active_library_transcriptions is not None:
            entry = state.active_library_transcriptions.get(library_audio_id)
            if entry is not None:
                entry['stage'] = 'diarizing'
                entry['progress'] = 1.0
        unique_speaker_labels: list[str] = []
        diar_embeddings: dict[str, list[float]] = {}
        if should_diarize:
            try:
                from app.services.diarization_service import (
                    DiarizationService, run_diarization_for_audio,
                )
                svc = DiarizationService.get_instance()
                if not svc.is_available():
                    logger.info(
                        "Diarization вимкнено: %s — пропускаю",
                        svc.unavailability_reason(),
                    )
                else:
                    rec_manifest = None
                    if source_type == 'recording':
                        sid = request.form.get('recording_session_id')
                        if sid and state.recording_service is not None:
                            try:
                                rec_manifest = state.recording_service.store.read(sid)
                            except Exception:
                                logger.debug(
                                    "recording manifest read failed for diarization sid=%s "
                                    "(continuing without manifest)", sid, exc_info=True,
                                )
                                rec_manifest = None
                    diar_t0 = time.time()
                    enriched_segs, unique_speaker_labels, diar_time, diar_embeddings = (
                        run_diarization_for_audio(
                            filepath,
                            result['segments'],
                            source_type=source_type,
                            recording_manifest=rec_manifest,
                            return_embeddings=True,  # Phase 10.6 voice fingerprinting
                        )
                    )
                    result['segments'] = enriched_segs
                    result['diarization_time'] = round(diar_time, 2)
                    result['speakers_detected'] = unique_speaker_labels
                    logger.info(
                        "Diarization: знайдено %d спікерів за %.1fs (wall=%.1fs, embeddings=%d)",
                        len(unique_speaker_labels), diar_time,
                        time.time() - diar_t0, len(diar_embeddings),
                    )
            except Exception as e:
                logger.warning(
                    "Diarization failed (продовжую без неї): %s", e, exc_info=True,
                )

        # Phase 14: напрямок (категорія) — з форми, опціонально
        _cat_raw = request.form.get('category_id')
        category_id = int(_cat_raw) if (_cat_raw and _cat_raw.isdigit()) else None

        with _get_db() as conn:
            c = conn.cursor()
            # Дедуп аудіо (Хвиля A): той самий текст, залитий удруге, лишається
            # записом (свої коментарі/файл/задачі), але позначається
            # duplicate_of → id оригіналу і НЕ індексується — інакше обидві
            # копії конкурують за слоти видачі однаковим вмістом.
            text_hash = dedup_audio.hash_for(source_type, result['text'])
            original = dedup_audio.find_original(conn, text_hash)
            duplicate_of = original['id'] if original else None
            if duplicate_of:
                logger.info("Дубль аудіо: текст збігається з #%s «%s» — запис "
                            "зберігаю, але не індексую", duplicate_of,
                            original.get('source_name'))
            c.execute('''INSERT INTO transcriptions
                         (source_type, source_name, source_url, youtube_id, youtube_title,
                          youtube_author, youtube_duration, youtube_thumbnail, file_path,
                          transcript_text, language, model_used, processing_time, segments,
                          category_id, content_hash, duplicate_of, title, description)
                         VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''', (
                source_type, source_name,
                youtube_info.get('original_url') if source_type == 'youtube' else None,
                youtube_info.get('video_id') if source_type == 'youtube' else None,
                youtube_info.get('title') if source_type == 'youtube' else None,
                youtube_info.get('author') if source_type == 'youtube' else None,
                youtube_info.get('duration') if source_type == 'youtube' else None,
                youtube_info.get('thumbnail') if source_type == 'youtube' else None,
                filepath, result['text'], result['language'], model_name,
                processing_time, json.dumps(result['segments']),
                category_id, text_hash, duplicate_of, record_title, record_description,
            ))
            transcription_id = c.lastrowid
            result['transcription_id'] = transcription_id
            if duplicate_of:
                result['duplicate_of'] = duplicate_of

            # Phase 10.3: створити transcription_speaker_map записи для виявлених
            # лейблів. 'self' автоматично мапиться на seeded speaker is_self=1 ('Ви'),
            # інші лейбли (SPEAKER_00, ...) залишаються з NULL speaker_id —
            # фронтенд покаже "Спікер N" і запропонує іменування.
            #
            # Phase 10.6: voice fingerprinting auto-match. Для кожного нового
            # SPEAKER_NN з embedding'ом порівнюємо cosine similarity до всіх
            # збережених speakers.embedding. Match >threshold → авто-link.
            if unique_speaker_labels:
                from app.services.diarization_service import (
                    embedding_to_blob, find_matching_speaker,
                )
                self_id_row = c.execute(
                    "SELECT id FROM speakers WHERE is_self = 1 LIMIT 1"
                ).fetchone()
                self_speaker_id = self_id_row[0] if self_id_row else None

                # Завантажуємо всі saved speakers з embeddings одним запитом
                saved_speakers = c.execute(
                    "SELECT id, embedding FROM speakers WHERE embedding IS NOT NULL AND is_self = 0"
                ).fetchall()
                saved_list = [(r[0], r[1]) for r in saved_speakers]

                auto_matched = 0
                for label in unique_speaker_labels:
                    embedding = diar_embeddings.get(label)
                    embedding_blob = embedding_to_blob(embedding) if embedding else None

                    if label == 'self':
                        sid_for_label = self_speaker_id
                    else:
                        # Спробувати auto-match за embedding
                        sid_for_label = None
                        if embedding:
                            match = find_matching_speaker(embedding, saved_list)
                            if match:
                                sid_for_label, sim = match
                                auto_matched += 1
                                logger.info(
                                    "Auto-match: %s → speaker_id=%d (cosine=%.3f)",
                                    label, sid_for_label, sim,
                                )

                    c.execute(
                        'INSERT OR IGNORE INTO transcription_speaker_map '
                        '(transcription_id, raw_label, speaker_id, embedding) '
                        'VALUES (?, ?, ?, ?)',
                        (transcription_id, label, sid_for_label, embedding_blob),
                    )
                if auto_matched:
                    logger.info(
                        "Voice fingerprinting: %d/%d спікерів авто-визначено",
                        auto_matched, len(unique_speaker_labels),
                    )
            conn.commit()
            # Збагачуємо response повним speaker payload — щоб UI міг одразу
            # рендерити chips без додаткового GET /api/history/<id>.
            if unique_speaker_labels:
                speaker_rows = c.execute('''
                    SELECT m.raw_label, m.speaker_id, s.name, s.color, s.is_self
                    FROM transcription_speaker_map m
                    LEFT JOIN speakers s ON s.id = m.speaker_id
                    WHERE m.transcription_id = ?
                ''', (transcription_id,)).fetchall()
                result['speakers'] = [
                    {
                        'raw_label': r['raw_label'],
                        'speaker_id': r['speaker_id'],
                        'name': r['name'],
                        'color': r['color'],
                        'is_self': bool(r['is_self']) if r['is_self'] is not None else False,
                    }
                    for r in speaker_rows
                ]

        state.metrics.inc("whisper_transcriptions_total", model=model_name, source=source_type)
        state.metrics.observe_duration("whisper_transcription_duration_seconds", processing_time, model=model_name)
        if unique_speaker_labels:
            state.metrics.inc(
                "whisper_diarizations_total", source=source_type,
            )

        # Phase 13: авто-збагачення (Meeting Memory). Фоновий job — НЕ блокує
        # відповідь. Витягує сутності (люди/проєкти/орг), summary, action items
        # одним Claude-викликом і наповнює наскрізний граф для RAG-пошуку.
        # Авто-вимикається якщо немає ANTHROPIC_API_KEY (enrichment.is_available()).
        # Дубль не збагачуємо і не чанкуємо: чанки оригіналу вже є, а копія
        # лише дублювала б їх у видачі (і платила б за Claude-картку вдруге).
        try:
            from app.services import enrichment
            if duplicate_of:
                result['enrichment'] = {"status": "skipped_duplicate"}
            elif enrichment.any_available() and state.job_queue is not None:
                _db_path = current_app.config['DATABASE']
                _enrich_ch = f"enrich_{transcription_id}"

                def _enrich_job(job, _tid=transcription_id, _db=_db_path, _ch=_enrich_ch):
                    try:
                        state.sse_broker.publish(_ch, "enrich", {"status": "running", "transcription_id": _tid})
                    except Exception:
                        logger.debug("SSE enrich 'running' publish failed for %s (best-effort)", _ch, exc_info=True)
                    res = enrichment.enrich_transcription(_db, _tid)
                    try:
                        state.sse_broker.publish(_ch, "complete", {"status": "completed", **res})
                    except Exception:
                        logger.debug("SSE enrich 'complete' publish failed for %s (best-effort)", _ch, exc_info=True)
                    return res

                state.job_queue.submit(
                    "enrichment", _enrich_job,
                    meta={"transcription_id": transcription_id},
                )
                result['enrichment'] = {"status": "queued"}
        except Exception as e:
            logger.warning("Enrichment job не поставлено (продовжую): %s", e)

        # Phase 19 (Крок 7): прилінкувати копілот-сесію запису до створеного
        # транскрипту → таймлайн ко-пілота відкривається зі сторінки транскрипту.
        if source_type == 'recording':
            _rec_sid = request.form.get('recording_session_id') or library_recording_sid
            if _rec_sid and state.copilot_service is not None:
                try:
                    _cs = state.copilot_service.get_by_recording(_rec_sid)
                    if _cs:
                        state.copilot_service.link_transcription(_cs['id'], transcription_id)
                except Exception as e:
                    logger.warning("copilot link to transcription failed: %s", e)

            # Волна 4: коментарі, написані ПІД ЧАС дзвінка, висіли на сесії
            # запису (транскрипту тоді ще не існувало). Тепер він є — і саме
            # тут вони мусять переїхати: інакше все, що власник надиктував по
            # ходу розмови, лишилось би на id, якого не видно з жодного екрана
            # архіву. anchor_time переїжджає як є — відлік у сесії й у
            # транскрипті той самий (секунди від початку запису).
            if _rec_sid:
                try:
                    from app.services import comments as comments_svc
                    moved = comments_svc.retarget(
                        current_app.config['DATABASE'],
                        'recording_session', _rec_sid,
                        'transcription', transcription_id)
                    if moved:
                        result['comments_moved'] = moved
                        logger.info("[comments] сесія %s → tx=%s: перецеплено %d",
                                    _rec_sid, transcription_id, moved)
                except Exception as e:
                    # Коментарі лишаються на сесії й не губляться; лікує
                    # повторний запуск або ручний retarget.
                    logger.warning("comments retarget failed: %s", e)

        # Phase 12.11 fix: НЕ видаляємо uploaded file після transcribe —
        # потрібен для audio player у transcript view (Phase 12.3+).
        # File залишається у UPLOAD_FOLDER; видаляється тільки при DELETE
        # transcription (handled у delete endpoints).
        if source_type == 'youtube' and download_id:
            state.add_log(download_id, "complete", "Процесс завершен успешно!", 100, "completed")
        return jsonify(result)

    except Exception as e:
        logger.error(f"Помилка транскрибування: {e}", exc_info=True)
        if source_type == 'youtube' and download_id:
            state.add_log(download_id, "transcribe", f"Помилка: {e}", -1, "error")
        # При помилці uploaded file видаляємо — у DB немає transcription'а на нього.
        if source_type == 'file' and filepath and os.path.exists(filepath):
            os.remove(filepath)
        return jsonify({"success": False, "error": "Помилка транскрибування аудіо"}), 500
    finally:
        if library_audio_id is not None and state.active_library_transcriptions is not None:
            state.active_library_transcriptions.pop(library_audio_id, None)


@transcription_bp.route('/api/transcribe/active', methods=['GET'])
def list_active_library_transcriptions():
    """Поточні /api/transcribe джоби стартовані з Audio Library.

    Frontend опитує цей endpoint щоб маркувати картки `.is-transcribing`
    (вижити після F5). Ключ — audio_download_id, тому одна картка = один
    активний джоб максимум.
    """
    store = state.active_library_transcriptions or {}
    return jsonify({'active': list(store.values())})


# ===== History =====

@transcription_bp.route('/api/history', methods=['GET'])
def get_history():
    """Список історії з фільтрами + FTS5 пошук."""
    page = max(1, request.args.get('page', 1, type=int))
    per_page = max(1, min(100, request.args.get('per_page', 10, type=int)))
    source_filter = request.args.get('source_type')
    language_filter = request.args.get('language')
    search_query = request.args.get('search', '')
    speaker_id_filter = request.args.get('speaker_id', type=int)  # Phase 12.7
    # Phase 15D: '' | 'all' → усі; 'none' → IS NULL; число → category_id = N.
    category_raw = (request.args.get('category_id') or '').strip().lower()
    # Контекстні фасет-фільтри (зʼявляються в UI під конкретне джерело):
    # Telegram — відправник + чат/група/канал; YouTube — автор/канал. LIKE по
    # метаданих (не транскрипту), тому доповнюють, а не замінюють FTS-пошук.
    tg_sender_q = (request.args.get('tg_sender') or '').strip()
    tg_chat_q = (request.args.get('tg_chat') or '').strip()
    yt_author_q = (request.args.get('yt_author') or '').strip()

    with _get_db() as conn:
        c = conn.cursor()
        # T4.6: soft-deleted (undo-вікно) ніколи не показуємо у списках/пошуку —
        # обов'язковий системний фільтр, не залежить від query-параметрів.
        where_clauses = ['deleted_at IS NULL']
        params = []
        fts_query = None  # Phase 12.2: lifted з if-block для use later у snippet
        if source_filter:
            where_clauses.append('source_type = ?'); params.append(source_filter)
        if language_filter:
            where_clauses.append('language = ?'); params.append(language_filter)
        if category_raw and category_raw != 'all':
            if category_raw == 'none':
                where_clauses.append('category_id IS NULL')
            elif category_raw.isdigit():
                where_clauses.append('category_id = ?'); params.append(int(category_raw))
        if speaker_id_filter:
            # Phase 12.7: фільтр по speaker_id — JOIN з map'ом, distinct щоб не
            # дублювати при кількох raw_label з тим же speaker_id (рідкісно, але буває).
            where_clauses.append(
                'id IN (SELECT DISTINCT transcription_id FROM transcription_speaker_map WHERE speaker_id = ?)'
            )
            params.append(speaker_id_filter)
        if tg_sender_q:
            where_clauses.append('tg_sender LIKE ?'); params.append(f'%{tg_sender_q}%')
        if tg_chat_q:
            where_clauses.append('tg_chat_title LIKE ?'); params.append(f'%{tg_chat_q}%')
        if yt_author_q:
            where_clauses.append('youtube_author LIKE ?'); params.append(f'%{yt_author_q}%')
        if search_query:
            fts_query = sanitize_fts_query(search_query)
            if fts_query:
                # Пошук по Бібліотеці бачить і КОМЕНТАРІ до запису, не лише його
                # текст (Волна 2.5). Інакше виходила дірка: власник пише
                # «Барселона» уточненням до дзвінка, де це слово не звучало, —
                # семантичний пошук той коментар знаходить, а Бібліотека, у якій
                # запис і треба відкрити, каже «нічого не знайдено».
                #
                # Через окремий FTS коментарів, а не LIKE: та сама токенізація,
                # що й у транскриптів, тож «бюджету» знаходить «бюджет».
                # Відсутність таблиці (стара БД) не має ламати пошук — тому
                # гілка перевіряє її наявність, а не покладається на except.
                cond = ('id IN (SELECT rowid FROM transcriptions_fts '
                        'WHERE transcriptions_fts MATCH ?)')
                params.append(fts_query)
                has_cm = bool(c.execute(
                    "SELECT 1 FROM sqlite_master WHERE type='table' "
                    "AND name='comment_chunks_fts'").fetchone())
                if has_cm:
                    cond = ('(' + cond + ' OR id IN ('
                            'SELECT c2.target_id FROM comments c2 '
                            'JOIN comment_chunks cc ON cc.comment_id = c2.id '
                            'JOIN comment_chunks_fts f ON f.rowid = cc.id '
                            "WHERE c2.target_type = 'transcription' "
                            'AND c2.deleted_at IS NULL '
                            'AND comment_chunks_fts MATCH ?))')
                    params.append(fts_query)
                where_clauses.append(cond)
            else:
                where_clauses.append('1 = 0')

        where_sql = 'WHERE ' + ' AND '.join(where_clauses) if where_clauses else ''
        count_sql = f'SELECT COUNT(*) FROM transcriptions {where_sql}'
        total = c.execute(count_sql, params).fetchone()[0]

        offset = (page - 1) * per_page
        # Phase 12.2: коли search активний, використовуємо FTS5 snippet()
        # для виділення фрагмента навколо matche'у. Інакше — звичайний preview
        # перших 200 символів.
        # Спільний набір колонок для обох гілок. Список у Бібліотеці рендерить
        # рядок ПО ТИПУ джерела (телеграм-повідомлення ≠ дзвінок ≠ документ), і
        # кожна форма потребує своїх полів: TG — чат/відправник/лінк, дзвінок —
        # тривалість/спікери/скільки задач, документ — сторінки/ім'я файлу.
        # Два корельовані підзапити на сторінку в 20 рядків — дешево.
        # `title`/`description` (v44) — поруч із `source_name`, не замість нього:
        # source_name лишається провенансом, display_name рахує record_meta.
        cols = '''t.id, t.created_at, t.source_type, t.source_name, t.title, t.description,
                       t.youtube_thumbnail,
                       t.language, t.model_used, t.processing_time, t.category_id, t.doc_type,
                       t.tg_chat_title, t.tg_sender, t.tg_link,
                       t.youtube_duration, t.youtube_author,
                       t.page_count, t.original_filename, t.meeting_date, t.enriched_at,
                       t.duplicate_of,
                       (SELECT COUNT(*) FROM action_items ai
                         WHERE ai.transcription_id = t.id
                           AND ai.status = 'open' AND ai.dup_of IS NULL) AS open_tasks,
                       (SELECT COUNT(DISTINCT m.speaker_id) FROM transcription_speaker_map m
                         WHERE m.transcription_id = t.id) AS speaker_count'''
        if fts_query:
            select_sql = f'''
                SELECT {cols},
                       (SELECT snippet(transcriptions_fts, 0, '«MARK»', '«/MARK»', '…', 16)
                        FROM transcriptions_fts
                        WHERE transcriptions_fts MATCH ? AND rowid = t.id) as transcript_preview
                FROM transcriptions t
                {where_sql.replace('id IN', 't.id IN')}
                ORDER BY t.created_at DESC
                LIMIT ?
                OFFSET ?
            '''
            transcriptions = c.execute(select_sql, [fts_query] + params + [per_page, offset]).fetchall()
        else:
            # Аліас `t` обов'язковий і тут: без нього `ai.transcription_id = id`
            # у підзапиті звʼязався б з action_items.id, а не з транскриптом.
            # Некваліфіковані імена у where_sql далі резолвляться в `t`.
            select_sql = f'''
                SELECT {cols},
                       SUBSTR(t.transcript_text, 1, 200) as transcript_preview
                FROM transcriptions t
                {where_sql}
                ORDER BY t.created_at DESC
                LIMIT ?
                OFFSET ?
            '''
            transcriptions = c.execute(select_sql, params + [per_page, offset]).fetchall()

    items = []
    for t in transcriptions:
        item = dict(t)
        item['display_name'] = record_meta.display_name(item)
        items.append(item)

    return jsonify({
        'transcriptions': items,
        'total': total,
        'page': page,
        'per_page': per_page,
        'total_pages': (total + per_page - 1) // per_page if total > 0 else 0,
    })


@transcription_bp.route('/api/history/<int:transcription_id>', methods=['GET'])
def get_transcription(transcription_id):
    """Один транскрипт зі всіма даними."""
    with _get_db() as conn:
        transcription = conn.execute(
            'SELECT * FROM transcriptions WHERE id = ? AND deleted_at IS NULL',
            (transcription_id,),
        ).fetchone()
        if not transcription:
            return jsonify({"success": False, "error": "Транскрибування не знайдено"}), 404
        # Phase 10.3: speaker map (raw_label → speaker info) для UI chip-рендеру.
        speaker_rows = conn.execute('''
            SELECT m.raw_label, m.speaker_id, s.name, s.color, s.is_self
            FROM transcription_speaker_map m
            LEFT JOIN speakers s ON s.id = m.speaker_id
            WHERE m.transcription_id = ?
        ''', (transcription_id,)).fetchall()

        # Зв'язки для нової сторінки транскрипту: сутності та action items
        # цього запису. Порожні списки, якщо enrichment ще не відпрацював.
        # Additive — наявні споживачі ігнорують нові ключі.
        entity_rows = conn.execute('''
            SELECT e.id, e.type, e.canonical_name, e.role,
                   me.mention_count, me.salience, me.role_in_meeting
            FROM meeting_entities me
            JOIN entities e ON e.id = me.entity_id
            WHERE me.transcription_id = ?
            ORDER BY me.salience DESC, me.mention_count DESC
        ''', (transcription_id,)).fetchall()
        action_item_rows = conn.execute('''
            SELECT id, task, owner_name, owner_entity_id, due, status, created_at
            FROM action_items
            WHERE transcription_id = ?
            ORDER BY (status = 'open') DESC, created_at ASC
        ''', (transcription_id,)).fetchall()
    result = dict(transcription)
    result['segments'] = json.loads(result['segments']) if result['segments'] else []
    result['speakers'] = [
        {
            'raw_label': r['raw_label'],
            'speaker_id': r['speaker_id'],
            'name': r['name'],
            'color': r['color'],
            'is_self': bool(r['is_self']) if r['is_self'] is not None else False,
        }
        for r in speaker_rows
    ]
    result['entities'] = [dict(r) for r in entity_rows]
    result['action_items'] = [dict(r) for r in action_item_rows]
    # title/description приходять із SELECT * (v44); display_name — похідне поле.
    result['display_name'] = record_meta.display_name(result)
    return jsonify(result)


@transcription_bp.route('/api/history/<int:transcription_id>', methods=['PATCH'])
def patch_transcription_meta(transcription_id):
    """Змінити власну назву/опис запису після збереження (спека editable-title-description).

    JSON ``{"title"?: str|null, "description"?: str|null}`` — міняються лише
    передані ключі. Порожній рядок означає «прибрати» (``NULL``), а НЕ копію
    ``source_name``: провенанс запису лишається недоторканим у будь-якому разі.
    """
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        return jsonify({"success": False, "error": "Очікується JSON-обʼєкт"}), 400
    if 'title' not in payload and 'description' not in payload:
        return jsonify({"success": False,
                        "error": "Потрібне хоча б одне поле: title або description"}), 400

    kwargs = {}
    if 'title' in payload:
        kwargs['title'] = payload['title']
    if 'description' in payload:
        kwargs['description'] = payload['description']

    try:
        with _get_db() as conn:
            record = record_meta.update_meta(conn, transcription_id, **kwargs)
    except ValueError as e:
        # Контрольована помилка нормалізації (довжина) — її текст писався для людини.
        return jsonify({"success": False, "error": str(e)}), 400

    if record is None:
        return jsonify({"success": False, "error": "Транскрибування не знайдено"}), 404

    changed = record.pop('changed', False)
    # db_path явно від запиту (той самий, яким відкрито зʼєднання вище): без
    # нього re-embed пішов би в Config.DATABASE і переписав чанки чужої БД.
    record_meta.after_meta_update(transcription_id, changed=changed,
                                  db_path=current_app.config['DATABASE'])
    return jsonify({"success": True, "record": record})


@transcription_bp.route('/api/history/<int:transcription_id>', methods=['DELETE'])
def delete_transcription(transcription_id):
    """Soft-delete одного транскрипту (T4.6, REMEDIATION_PLAN Волна 2, Варіант A).

    Ставить ``deleted_at`` замість фізичного видалення — файл (youtube/file
    джерело) НЕ стирається одразу, лишається доступним для
    ``POST /api/history/<id>/restore`` протягом grace-періоду
    (``RECALL_SOFTDELETE_GRACE_DAYS``, дефолт 7д; фізично прибирає
    ``app.services.retention.purge_soft_deleted`` на старті сервера).
    Bulk-delete (``/api/history/bulk_delete``) лишається фізичним — undo
    свідомо підтримується лише для одиничного delete.
    """
    try:
        with _get_db() as conn:
            row = conn.execute(
                'SELECT id FROM transcriptions WHERE id = ? AND deleted_at IS NULL',
                (transcription_id,),
            ).fetchone()
            if not row:
                return jsonify({"success": False, "error": "Транскрибування не знайдено"}), 404
            conn.execute(
                'UPDATE transcriptions SET deleted_at = ? WHERE id = ?',
                (time.time(), transcription_id),
            )
            conn.commit()
        return jsonify({"success": True, "id": transcription_id})
    except Exception as e:
        logger.error(f"Помилка видалення транскрибування: {e}")
        return jsonify({"success": False, "error": "Помилка видалення транскрибування"}), 500


@transcription_bp.route('/api/history/<int:transcription_id>/restore', methods=['POST'])
def restore_transcription(transcription_id):
    """Скасувати soft-delete (undo). 404 якщо запису нема або він не видалений."""
    try:
        with _get_db() as conn:
            row = conn.execute(
                'SELECT id FROM transcriptions WHERE id = ? AND deleted_at IS NOT NULL',
                (transcription_id,),
            ).fetchone()
            if not row:
                return jsonify({"success": False, "error": "Немає що відновлювати"}), 404
            conn.execute(
                'UPDATE transcriptions SET deleted_at = NULL WHERE id = ?',
                (transcription_id,),
            )
            conn.commit()
        return jsonify({"success": True, "id": transcription_id})
    except Exception as e:
        logger.error(f"Помилка відновлення транскрибування: {e}")
        return jsonify({"success": False, "error": "Помилка відновлення транскрибування"}), 500


@transcription_bp.route('/api/history/bulk_delete', methods=['POST'])
def bulk_delete_transcriptions():
    """Масове видалення."""
    try:
        data = request.json
        ids = data.get('ids', [])
        if not ids:
            return jsonify({"success": False, "error": "Не вказано ID для видалення"}), 400
        try:
            ids = [int(id) for id in ids]
        except ValueError:
            return jsonify({"success": False, "error": "Невірний формат ID"}), 400

        with _get_db() as conn:
            placeholders = ','.join('?' * len(ids))
            files_info = conn.execute(
                f'SELECT id, file_path, source_type FROM transcriptions WHERE id IN ({placeholders})',
                ids,
            ).fetchall()
            # Phase 10.3: чистимо speaker map перед видаленням transcriptions
            conn.execute(
                f'DELETE FROM transcription_speaker_map WHERE transcription_id IN ({placeholders})',
                ids,
            )
            conn.execute(f'DELETE FROM transcriptions WHERE id IN ({placeholders})', ids)
            conn.commit()

        deleted_files = 0
        # Phase 12.11: file + youtube → видаляємо. recording → ні.
        for file_info in files_info:
            if file_info['source_type'] in ('youtube', 'file') and file_info['file_path']:
                try:
                    if os.path.exists(file_info['file_path']):
                        os.remove(file_info['file_path'])
                        deleted_files += 1
                except Exception as e:
                    logger.error(f"Помилка видалення файлу {file_info['file_path']}: {e}")
        return jsonify({
            "message": f"Видалено {len(files_info)} транскрибувань",
            "deleted_files": deleted_files,
        })
    except Exception as e:
        logger.error(f"Помилка масового видалення: {e}")
        return jsonify({"success": False, "error": "Помилка масового видалення"}), 500


@transcription_bp.route('/api/history/bulk_export', methods=['POST'])
def bulk_export_transcriptions():
    """Масовий експорт (JSON/TXT/MD).

    Phase 10.9: для діаризованих транскриптів TXT і MD містять speaker
    префікси (Імʼя: текст). JSON включає resolved speakers map per item.
    """
    try:
        data = request.json
        ids = data.get('ids', [])
        export_format = data.get('format', 'json')
        if not ids:
            return jsonify({"success": False, "error": "Не вказано ID для експорту"}), 400
        try:
            ids = [int(id) for id in ids]
        except ValueError:
            return jsonify({"success": False, "error": "Невірний формат ID"}), 400

        with _get_db() as conn:
            placeholders = ','.join('?' * len(ids))
            transcriptions = conn.execute(
                f'SELECT * FROM transcriptions WHERE id IN ({placeholders})',
                ids,
            ).fetchall()
            if not transcriptions:
                return jsonify({"success": False, "error": "Транскрибування не знайдено"}), 404

            # Phase 10.9: підтягуємо speaker maps для всіх transcripts одним запитом
            speaker_maps_rows = conn.execute(f'''
                SELECT m.transcription_id, m.raw_label, s.name, s.is_self
                FROM transcription_speaker_map m
                LEFT JOIN speakers s ON s.id = m.speaker_id
                WHERE m.transcription_id IN ({placeholders})
            ''', ids).fetchall()

        # Build per-transcription speaker map: {tid: [{raw_label, name, is_self}]}
        speaker_maps_by_tid: dict[int, list[dict]] = {}
        for r in speaker_maps_rows:
            speaker_maps_by_tid.setdefault(r['transcription_id'], []).append({
                'raw_label': r['raw_label'],
                'name': r['name'],
                'is_self': bool(r['is_self']) if r['is_self'] is not None else False,
            })

        def _resolver(speakers_list):
            """Повертає функцію resolve(raw_label) → display name."""
            display = {s['raw_label']: s['name'] for s in speakers_list if s.get('name')}
            def _r(raw_label: str) -> str:
                if not raw_label:
                    return ''
                if raw_label in display:
                    return display[raw_label]
                if raw_label == 'self':
                    return 'Ви'
                if raw_label == 'SPEAKER_UNKNOWN':
                    return '?'
                m = re.match(r'^SPEAKER_(\d+)$', raw_label)
                return f'Спікер {int(m.group(1)) + 1}' if m else raw_label
            return _r

        def _fmt_ts(seconds: float) -> str:
            mins, secs = divmod(int(seconds), 60)
            hrs, mins = divmod(mins, 60)
            return f'{hrs:02d}:{mins:02d}:{secs:02d}' if hrs else f'{mins:02d}:{secs:02d}'

        export_data = []
        for t in transcriptions:
            item = dict(t)
            item['segments'] = json.loads(item['segments']) if item['segments'] else []
            item['speakers'] = speaker_maps_by_tid.get(item['id'], [])
            # Коментарі власника (Волна 5) — по одному запиту на запис. Дешево
            # (індекс по target) і в масштабі bulk-експорту незначуще проти
            # витягування самих транскриптів.
            item['comments'] = _export_comments(item['id'])
            item['display_name'] = record_meta.display_name(item)
            export_data.append(item)

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        if export_format == 'json':
            export_path = os.path.join(current_app.config['TRANSCRIPTS_FOLDER'], f"bulk_export_{timestamp}.json")
            with open(export_path, 'w', encoding='utf-8') as f:
                json.dump(export_data, f, ensure_ascii=False, indent=2, default=str)
        elif export_format == 'txt':
            export_path = os.path.join(current_app.config['TRANSCRIPTS_FOLDER'], f"bulk_export_{timestamp}.txt")
            with open(export_path, 'w', encoding='utf-8') as f:
                for item in export_data:
                    f.write(f"\n{'='*60}\n")
                    f.write(f"Джерело: {item['display_name']}\n")
                    f.write(f"Дата: {item['created_at']}\n")
                    f.write(f"Мова: {item['language']}\n")
                    f.write(f"{'='*60}\n\n")
                    f.write(_comments_text(item.get('comments') or []))
                    segs = item['segments']
                    has_sp = any(s.get('speaker') for s in segs)
                    if has_sp and segs:
                        resolve = _resolver(item['speakers'])
                        prev = None
                        for seg in segs:
                            sp = resolve(seg.get('speaker', ''))
                            text = (seg.get('text') or '').strip()
                            if not text:
                                continue
                            if sp != prev:
                                if prev is not None:
                                    f.write('\n\n')
                                f.write(f'{sp}: {text}')
                                prev = sp
                            else:
                                f.write(f' {text}')
                        f.write('\n')
                    else:
                        f.write(item.get('transcript_text', ''))
                    f.write("\n\n")
        elif export_format == 'docx':
            try:
                from docx import Document
                from docx.shared import Pt
            except ImportError:
                return jsonify({
                    "success": False,
                    "error": "python-docx не встановлений",
                }), 500
            export_path = os.path.join(current_app.config['TRANSCRIPTS_FOLDER'], f"bulk_export_{timestamp}.docx")
            doc = Document()
            for idx, item in enumerate(export_data):
                if idx > 0:
                    doc.add_page_break()
                doc.add_heading(item['display_name'] or 'Транскрипт', level=1)
                meta = doc.add_paragraph()
                meta.add_run(f'Дата: {item["created_at"]}   ·   Мова: {item["language"]}').italic = True
                for c in (item.get('comments') or []):
                    p = doc.add_paragraph()
                    ts = _cm_ts(c.get('anchor_time'))
                    p.add_run(f'{_CM_LABEL.get(c.get("kind"), "НОТАТКА")}'
                              f'{" · " + ts if ts else ""}: ').bold = True
                    p.add_run(c.get('body') or '')
                    p.paragraph_format.space_after = Pt(6)
                segs = item['segments']
                has_sp = any(s.get('speaker') for s in segs)
                if has_sp and segs:
                    from collections import Counter
                    resolve = _resolver(item['speakers'])
                    spk_counts = Counter(
                        resolve(s.get('speaker', ''))
                        for s in segs if (s.get('text') or '').strip()
                    )
                    sp_p = doc.add_paragraph()
                    sp_p.add_run('Спікери: ').bold = True
                    sp_p.add_run(', '.join(f'{n} ({c})' for n, c in spk_counts.most_common()))
                    doc.add_paragraph()

                    prev = None
                    turn: list[str] = []
                    turn_start = 0.0
                    def _flush():
                        if prev and turn:
                            h = doc.add_heading(level=3)
                            h.add_run(prev).bold = True
                            h.add_run(f'   ·   {_fmt_ts(turn_start)}').italic = True
                            p = doc.add_paragraph(' '.join(turn))
                            p.paragraph_format.space_after = Pt(8)
                    for seg in segs:
                        sp = resolve(seg.get('speaker', ''))
                        text = (seg.get('text') or '').strip()
                        if not text:
                            continue
                        if sp != prev:
                            _flush()
                            prev = sp
                            turn = [text]
                            turn_start = float(seg.get('start', 0.0))
                        else:
                            turn.append(text)
                    _flush()
                else:
                    doc.add_paragraph(item.get('transcript_text', ''))
            doc.save(export_path)
        elif export_format == 'md':
            export_path = os.path.join(current_app.config['TRANSCRIPTS_FOLDER'], f"bulk_export_{timestamp}.md")
            with open(export_path, 'w', encoding='utf-8') as f:
                for item in export_data:
                    f.write(f'# {item["display_name"]}\n\n')
                    f.write(f'**Дата:** {item["created_at"]} · **Мова:** {item["language"]}\n\n')
                    f.write(_comments_md(item.get('comments') or []))
                    segs = item['segments']
                    has_sp = any(s.get('speaker') for s in segs)
                    if has_sp and segs:
                        from collections import Counter
                        resolve = _resolver(item['speakers'])
                        spk_counts = Counter(
                            resolve(s.get('speaker', ''))
                            for s in segs if (s.get('text') or '').strip()
                        )
                        spk_line = ', '.join(f'**{n}** ({c})' for n, c in spk_counts.most_common())
                        f.write(f'**Спікери:** {spk_line}\n\n---\n\n')

                        prev = None
                        turn: list[str] = []
                        turn_start = 0.0
                        def _flush():
                            if prev and turn:
                                f.write(f'### {prev} · `{_fmt_ts(turn_start)}`\n\n')
                                f.write(' '.join(turn))
                                f.write('\n\n')
                        for seg in segs:
                            sp = resolve(seg.get('speaker', ''))
                            text = (seg.get('text') or '').strip()
                            if not text:
                                continue
                            if sp != prev:
                                _flush()
                                prev = sp
                                turn = [text]
                                turn_start = float(seg.get('start', 0.0))
                            else:
                                turn.append(text)
                        _flush()
                    else:
                        f.write('---\n\n')
                        f.write(item.get('transcript_text', ''))
                        f.write('\n\n')
        else:
            return jsonify({"success": False, "error": "Непідтримуваний формат експорту"}), 400

        response = send_file(export_path, as_attachment=True)

        @response.call_on_close
        def _cleanup_export():
            try:
                if os.path.exists(export_path):
                    os.remove(export_path)
            except Exception:
                logger.debug("export cleanup failed for %s (best-effort)", export_path, exc_info=True)
        return response
    except Exception as e:
        logger.error(f"Помилка масового експорту: {e}")
        return jsonify({"success": False, "error": "Помилка масового експорту"}), 500


@transcription_bp.route('/api/transcription/<int:transcription_id>/audio', methods=['GET'])
def get_transcription_audio(transcription_id):
    """Phase 12.3: streamити вихідний аудіо-файл транскрипції для player'а.

    Range-headers підтримує send_file автоматично — браузер може seek'ати
    у середину файлу для click-on-segment-to-play.
    """
    with _get_db() as conn:
        row = conn.execute(
            'SELECT file_path, source_type FROM transcriptions WHERE id = ? AND deleted_at IS NULL',
            (transcription_id,),
        ).fetchone()
    if not row:
        return jsonify({"success": False, "error": "Транскрипт не знайдено"}), 404
    file_path = row['file_path']
    if not file_path or not os.path.isfile(file_path):
        return jsonify({"success": False, "error": "Аудіо-файл недоступний"}), 404
    # Path-traversal protection: file_path має бути у відомих директоріях
    allowed_roots = [
        current_app.config.get('UPLOAD_FOLDER', ''),
        os.path.join(current_app.config.get('BASE_DIR', '.'), 'youtube_downloads'),
        os.path.join(current_app.config.get('BASE_DIR', '.'), 'recordings'),
    ]
    abs_path = safe_path_within_any(allowed_roots, file_path)
    if abs_path is None:
        logger.warning(f"Audio request blocked (path outside allowed roots): {file_path}")
        return jsonify({"success": False, "error": "Access denied"}), 403
    return send_file(abs_path, conditional=True)


# ============================================================
# Phase 12.9: Manual segment merge / split
# ============================================================

def _save_segments(conn, transcription_id: int, segments: list[dict]) -> None:
    """Зберігає segments JSON та оновлює transcript_text.

    transcript_text = ' '.join(seg.text) — використовується у FTS5 і
    /api/history list. Тригери на FTS оновлюються автоматично.
    Якщо є polished_text — він НЕ перезаписується (зберігає Claude версію).
    """
    transcript_text = ' '.join(s.get('text', '').strip() for s in segments).strip()
    conn.execute(
        'UPDATE transcriptions SET segments = ?, transcript_text = ? WHERE id = ?',
        (json.dumps(segments, ensure_ascii=False), transcript_text, transcription_id),
    )
    # Інвалідуємо кеш summary (текст змінився).
    conn.execute(
        'UPDATE transcriptions SET summary_json = NULL, summary_at = NULL, '
        'summary_model = NULL WHERE id = ?',
        (transcription_id,),
    )


@transcription_bp.route('/api/transcription/<int:transcription_id>/segments/merge', methods=['POST'])
def merge_segments(transcription_id):
    """Об'єднати кілька consecutive сегментів у один.

    Body: {"indices": [i, i+1, ...]} — мають бути послідовні числа.
    Speaker беремо з першого, time = [first.start, last.end],
    text = ' '.join всіх text'ів trimmed, words = concat (якщо є).
    """
    data = request.get_json(silent=True) or {}
    indices = data.get('indices')
    if not isinstance(indices, list) or len(indices) < 2:
        return jsonify({'success': False, 'error': 'indices мусить бути списком ≥2 елементів'}), 400
    try:
        indices = sorted(set(int(x) for x in indices))
    except (TypeError, ValueError):
        return jsonify({'success': False, 'error': 'indices — лише числа'}), 400
    # Перевірка послідовності
    if any(indices[k+1] - indices[k] != 1 for k in range(len(indices) - 1)):
        return jsonify({'success': False, 'error': 'indices мусять бути послідовні'}), 400

    with _get_db() as conn:
        row = conn.execute(
            'SELECT segments FROM transcriptions WHERE id = ?',
            (transcription_id,),
        ).fetchone()
        if not row:
            return jsonify({'success': False, 'error': 'Транскрипт не знайдено'}), 404
        segments = json.loads(row['segments']) if row['segments'] else []
        if not segments:
            return jsonify({'success': False, 'error': 'Сегментів немає'}), 400
        if indices[0] < 0 or indices[-1] >= len(segments):
            return jsonify({'success': False, 'error': 'indices виходять за межі'}), 400

        first, last = segments[indices[0]], segments[indices[-1]]
        merged = {
            'id': first.get('id', indices[0]),
            'start': first.get('start', 0.0),
            'end': last.get('end', last.get('start', 0.0)),
            'text': ' '.join((segments[i].get('text') or '').strip() for i in indices).strip(),
            'speaker': first.get('speaker'),
        }
        # Concat words якщо є
        words: list = []
        for i in indices:
            w = segments[i].get('words')
            if isinstance(w, list):
                words.extend(w)
        if words:
            merged['words'] = words

        new_segments = segments[:indices[0]] + [merged] + segments[indices[-1] + 1:]
        # Re-id послідовно
        for k, s in enumerate(new_segments):
            s['id'] = k
        _save_segments(conn, transcription_id, new_segments)
        conn.commit()

    logger.info(f"[segments/merge] tx={transcription_id}: merged {len(indices)} segs at {indices[0]}")
    return jsonify({'success': True, 'segments': new_segments, 'merged_index': indices[0]})


@transcription_bp.route('/api/transcription/<int:transcription_id>/segments/split', methods=['POST'])
def split_segment(transcription_id):
    """Розділити segments[index] у точці split_at_seconds.

    Body: {
        "index": int,
        "split_at_seconds": float,    // абсолютний час (між seg.start і seg.end)
        "right_speaker": str|null     // optional — змінити speaker правої половини
    }

    Текст ділиться пропорційно: якщо є words[] — по найближчому word
    boundary, інакше по character ratio (split_at - start) / duration.
    """
    data = request.get_json(silent=True) or {}
    try:
        index = int(data.get('index'))
        split_at = float(data.get('split_at_seconds'))
    except (TypeError, ValueError):
        return jsonify({'success': False, 'error': 'index/split_at_seconds невірні'}), 400
    right_speaker = data.get('right_speaker')  # None = lefthand speaker

    with _get_db() as conn:
        row = conn.execute(
            'SELECT segments FROM transcriptions WHERE id = ?',
            (transcription_id,),
        ).fetchone()
        if not row:
            return jsonify({'success': False, 'error': 'Транскрипт не знайдено'}), 404
        segments = json.loads(row['segments']) if row['segments'] else []
        if index < 0 or index >= len(segments):
            return jsonify({'success': False, 'error': 'index виходить за межі'}), 400

        seg = segments[index]
        start = float(seg.get('start', 0))
        end = float(seg.get('end', start))
        if not (start < split_at < end):
            return jsonify({
                'success': False,
                'error': f'split_at ({split_at:.2f}s) має бути всередині сегменту [{start:.2f}, {end:.2f}]',
            }), 400

        text = (seg.get('text') or '').strip()
        words = seg.get('words') if isinstance(seg.get('words'), list) else None

        left_text = ''
        right_text = ''
        left_words = None
        right_words = None
        if words:
            left_words = []
            right_words = []
            for w in words:
                w_start = float(w.get('start', start))
                if w_start < split_at:
                    left_words.append(w)
                else:
                    right_words.append(w)
            left_text = ' '.join((w.get('word') or w.get('text') or '').strip() for w in left_words).strip()
            right_text = ' '.join((w.get('word') or w.get('text') or '').strip() for w in right_words).strip()
        if not left_text and not right_text:
            # Fallback — character ratio
            ratio = (split_at - start) / max(0.001, end - start)
            cut = max(1, min(len(text) - 1, int(round(len(text) * ratio))))
            # Прив'язати до word boundary (whitespace)
            while cut > 0 and not text[cut].isspace():
                cut -= 1
            if cut == 0:
                cut = max(1, int(round(len(text) * ratio)))
            left_text = text[:cut].strip()
            right_text = text[cut:].strip()

        left_seg = {
            'id': index,
            'start': start,
            'end': split_at,
            'text': left_text,
            'speaker': seg.get('speaker'),
        }
        if left_words:
            left_seg['words'] = left_words
        right_seg = {
            'id': index + 1,
            'start': split_at,
            'end': end,
            'text': right_text,
            'speaker': right_speaker if right_speaker else seg.get('speaker'),
        }
        if right_words:
            right_seg['words'] = right_words

        new_segments = segments[:index] + [left_seg, right_seg] + segments[index + 1:]
        for k, s in enumerate(new_segments):
            s['id'] = k
        _save_segments(conn, transcription_id, new_segments)
        conn.commit()

    logger.info(f"[segments/split] tx={transcription_id}: split index {index} at {split_at:.2f}s")
    return jsonify({'success': True, 'segments': new_segments, 'left_index': index})


# ============================================================
# T7.7: shared helpers for the five near-identical Claude endpoints below
# (polish/summarize/sentiment/translate/topics) — each used to inline its own
# copy of "parse segments → load raw_label→name speaker map" and "call
# text_polishing.* → RuntimeError becomes 400, anything else logged + generic
# 500". Persistence/response shape genuinely differ per endpoint (see
# REMEDIATION_PLAN T7.7) and are intentionally NOT folded in here.
# ============================================================

def _load_speaker_context(
    conn, transcription_id: int, segments_json: str | None,
) -> tuple[list[dict], bool, dict[str, str]]:
    """Parse `segments` JSON and, if diarized, load the raw_label→name map."""
    segments = json.loads(segments_json) if segments_json else []
    has_speakers = any(s.get('speaker') for s in segments)
    speaker_map: dict[str, str] = {}
    if has_speakers:
        sm_rows = conn.execute('''
            SELECT m.raw_label, s.name
            FROM transcription_speaker_map m
            LEFT JOIN speakers s ON s.id = m.speaker_id
            WHERE m.transcription_id = ?
        ''', (transcription_id,)).fetchall()
        for r in sm_rows:
            if r['name']:
                speaker_map[r['raw_label']] = r['name']
    return segments, has_speakers, speaker_map


def _call_claude_endpoint(label: str, fn, *args, error_msg: str = "Помилка Claude API. Перевірте логи.", **kwargs):
    """Run a text_polishing.* call with the standard error handling.

    Returns (result, None) on success, or (None, (response, status)) on
    error — callers do ``result, err = _call_claude_endpoint(...); if err:
    return err``. error_msg is intentionally per-call — the five original
    endpoints don't all use the same wording, and this doesn't change that.
    """
    try:
        return fn(*args, **kwargs), None
    except RuntimeError as e:
        return None, (jsonify({"success": False, "error": str(e)}), 400)
    except Exception as e:
        logger.error(f"[{label}] error: {e}", exc_info=True)
        return None, (jsonify({"success": False, "error": error_msg}), 500)


@transcription_bp.route('/api/transcription/<int:transcription_id>/polish', methods=['POST'])
def polish_transcription(transcription_id):
    """Покращити транскрипт через Claude API. Кешує результат у БД.

    Phase 10.7: для diarized транскриптів (segments мають speaker field)
    автоматично використовується polish_diarized_transcript який зберігає
    speaker-розмітку у форматі 'Імʼя: текст'.
    """
    if not text_polishing.is_available():
        return jsonify({
            "success": False,
            "error": "ANTHROPIC_API_KEY не встановлено. Створіть .env у корені проекту з ключем.",
        }), 400

    data = request.get_json(silent=True) or {}
    model_override = data.get('model')
    force = bool(data.get('force', False))

    with _get_db() as conn:
        row = tx_repo.get_by_id(conn, transcription_id,
                                 columns=['id', 'transcript_text', 'polished_text', 'segments'])
        if not row:
            return jsonify({"success": False, "error": "Транскрибування не знайдено"}), 404

        if row['polished_text'] and not force:
            return jsonify({"success": True, "polished_text": row['polished_text'], "cached": True})

        # Phase 10.7: detect diarized transcript and load speaker map
        segments, has_speakers, speaker_map = _load_speaker_context(conn, transcription_id, row['segments'])

    raw_text = row['transcript_text'] or ''
    if not raw_text.strip() and not has_speakers:
        return jsonify({"success": False, "error": "Транскрипт порожній"}), 400

    if has_speakers:
        logger.info(
            "[polish] using diarized polish for transcription %d (%d segments, %d named speakers)",
            transcription_id, len(segments), len(speaker_map),
        )
        result, err = _call_claude_endpoint(
            "polish", text_polishing.polish_diarized_transcript, segments, speaker_map, model=model_override,
        )
    else:
        result, err = _call_claude_endpoint(
            "polish", text_polishing.polish_transcript, raw_text, model=model_override,
        )
    if err:
        return err

    polished = result["polished_text"]
    if not polished:
        return jsonify({"success": False, "error": "Claude повернув порожній текст"}), 500

    with _get_db() as conn:
        conn.execute(
            'UPDATE transcriptions SET polished_text = ?, polished_at = CURRENT_TIMESTAMP, polished_model = ? WHERE id = ?',
            (polished, result["model"], transcription_id),
        )
        conn.commit()

    state.metrics.inc("whisper_polish_total", model=result["model"])
    state.metrics.inc("whisper_polish_tokens_total", value=result["input_tokens"], model=result["model"], kind="input")
    state.metrics.inc("whisper_polish_tokens_total", value=result["output_tokens"], model=result["model"], kind="output")
    state.metrics.inc("whisper_polish_tokens_total", value=result["cache_read_tokens"], model=result["model"], kind="cache_read")

    return jsonify({
        "success": True,
        "polished_text": polished,
        "cached": False,
        "model": result["model"],
        "usage": {
            "input_tokens": result["input_tokens"],
            "output_tokens": result["output_tokens"],
            "cache_read_tokens": result["cache_read_tokens"],
            "cache_creation_tokens": result["cache_creation_tokens"],
        },
    })


@transcription_bp.route('/api/transcription/<int:transcription_id>/summarize', methods=['POST'])
def summarize_transcription(transcription_id):
    """Phase 12.5: згенерувати summary + key points + action items через Claude.

    Кешує JSON у summary_json. Body: {"force": bool, "model": str}.
    """
    if not text_polishing.is_available():
        return jsonify({
            "success": False,
            "error": "ANTHROPIC_API_KEY не встановлено. Створіть .env у корені проекту з ключем.",
        }), 400

    data = request.get_json(silent=True) or {}
    model_override = data.get('model')
    force = bool(data.get('force', False))

    with _get_db() as conn:
        row = tx_repo.get_by_id(conn, transcription_id, columns=[
            'id', 'transcript_text', 'polished_text', 'segments', 'summary_json',
        ])
        if not row:
            return jsonify({"success": False, "error": "Транскрибування не знайдено"}), 404

        if row['summary_json'] and not force:
            try:
                parsed = json.loads(row['summary_json'])
                return jsonify({"success": True, "cached": True, **parsed})
            except json.JSONDecodeError:
                pass  # broken cache → regen

        segments, has_speakers, speaker_map = _load_speaker_context(conn, transcription_id, row['segments'])

    raw_text = row['transcript_text'] or ''
    polished = row['polished_text'] or ''

    # Перевага: polished > diarized formatted > raw
    diarized_text = None
    if has_speakers and speaker_map:
        # Reuse polish formatter — формує 'Ім'я: текст\n\nІм'я: текст'.
        diarized_text = text_polishing._format_diarized_input(segments, speaker_map)

    body = polished if polished.strip() else (diarized_text or raw_text)
    if not body or not body.strip():
        return jsonify({"success": False, "error": "Транскрипт порожній"}), 400

    result, err = _call_claude_endpoint(
        "summarize", text_polishing.summarize_transcript,
        text=body if not diarized_text else raw_text,
        diarized_text=diarized_text,
        model=model_override,
    )
    if err:
        return err

    payload = {
        "summary": result["summary"],
        "key_points": result["key_points"],
        "action_items": result["action_items"],
    }

    with _get_db() as conn:
        conn.execute(
            'UPDATE transcriptions SET summary_json = ?, summary_at = CURRENT_TIMESTAMP, '
            'summary_model = ? WHERE id = ?',
            (json.dumps(payload, ensure_ascii=False), result["model"], transcription_id),
        )
        conn.commit()

    state.metrics.inc("whisper_summarize_total", model=result["model"])
    state.metrics.inc("whisper_summarize_tokens_total", value=result["input_tokens"], model=result["model"], kind="input")
    state.metrics.inc("whisper_summarize_tokens_total", value=result["output_tokens"], model=result["model"], kind="output")

    return jsonify({
        "success": True,
        "cached": False,
        "model": result["model"],
        "usage": {
            "input_tokens": result["input_tokens"],
            "output_tokens": result["output_tokens"],
            "cache_read_tokens": result["cache_read_tokens"],
        },
        **payload,
    })


@transcription_bp.route('/api/transcription/<int:transcription_id>/sentiment', methods=['POST'])
def sentiment_transcription(transcription_id):
    """Phase 12.25: per-speaker sentiment analysis. Потребує diarized transcript."""
    if not text_polishing.is_available():
        return jsonify({"success": False, "error": "ANTHROPIC_API_KEY не встановлено."}), 400

    data = request.get_json(silent=True) or {}
    force = bool(data.get('force', False))
    model_override = data.get('model')

    with _get_db() as conn:
        row = tx_repo.get_by_id(conn, transcription_id, columns=['id', 'segments', 'sentiment_json'])
        if not row:
            return jsonify({"success": False, "error": "Транскрибування не знайдено"}), 404

        if row['sentiment_json'] and not force:
            try:
                cached = json.loads(row['sentiment_json'])
                return jsonify({"success": True, "cached": True, **cached})
            except json.JSONDecodeError:
                pass

        segments, has_speakers, speaker_map = _load_speaker_context(conn, transcription_id, row['segments'])
        if not has_speakers:
            return jsonify({
                "success": False,
                "error": "Sentiment per speaker потребує діаризованого транскрипту",
            }), 400

    diarized_text = text_polishing._format_diarized_input(segments, speaker_map)
    if not diarized_text.strip():
        return jsonify({"success": False, "error": "Транскрипт порожній"}), 400

    result, err = _call_claude_endpoint(
        "sentiment", text_polishing.analyze_sentiment, diarized_text, model=model_override,
        error_msg="Помилка Claude API.",
    )
    if err:
        return err

    payload = {"speakers": result['speakers'], "model": result['model']}
    with _get_db() as conn:
        conn.execute(
            'UPDATE transcriptions SET sentiment_json = ? WHERE id = ?',
            (json.dumps(payload, ensure_ascii=False), transcription_id),
        )
        conn.commit()

    state.metrics.inc("whisper_sentiment_total", model=result["model"])

    return jsonify({"success": True, "cached": False, **payload})


@transcription_bp.route('/api/transcription/<int:transcription_id>/translate', methods=['POST'])
def translate_transcription(transcription_id):
    """Phase 12.11: переклад транскрипту через Claude.

    Body: {"target_lang": "en"|"uk"|"ru", "force": bool, "model": str}.
    Кешує у translations_json: {lang_code: {text, model, created_at}}.
    """
    if not text_polishing.is_available():
        return jsonify({
            "success": False,
            "error": "ANTHROPIC_API_KEY не встановлено.",
        }), 400

    data = request.get_json(silent=True) or {}
    target_lang = (data.get('target_lang') or '').strip().lower()
    if not target_lang or len(target_lang) > 5:
        return jsonify({"success": False, "error": "target_lang обов'язковий (e.g. 'en')"}), 400
    model_override = data.get('model')
    force = bool(data.get('force', False))

    with _get_db() as conn:
        row = tx_repo.get_by_id(conn, transcription_id, columns=[
            'id', 'transcript_text', 'polished_text', 'segments', 'language', 'translations_json',
        ])
        if not row:
            return jsonify({"success": False, "error": "Транскрибування не знайдено"}), 404

        if row['language'] and row['language'].lower() == target_lang:
            return jsonify({
                "success": False,
                "error": f"Транскрипт вже на мові '{target_lang}'",
            }), 400

        translations = {}
        if row['translations_json']:
            try:
                translations = json.loads(row['translations_json'])
            except json.JSONDecodeError:
                translations = {}

        if target_lang in translations and not force:
            return jsonify({
                "success": True,
                "cached": True,
                "translated_text": translations[target_lang].get('text', ''),
                "target_lang": target_lang,
                "model": translations[target_lang].get('model', ''),
            })

        segments, has_speakers, speaker_map = _load_speaker_context(conn, transcription_id, row['segments'])

    raw_text = row['transcript_text'] or ''
    polished = row['polished_text'] or ''
    diarized_text = None
    if has_speakers and speaker_map:
        diarized_text = text_polishing._format_diarized_input(segments, speaker_map)
    body_text = polished if polished.strip() else raw_text

    if not body_text.strip() and not diarized_text:
        return jsonify({"success": False, "error": "Транскрипт порожній"}), 400

    result, err = _call_claude_endpoint(
        "translate", text_polishing.translate_transcript,
        text=body_text, target_lang=target_lang, diarized_text=diarized_text, model=model_override,
    )
    if err:
        return err

    translated = result['translated_text']
    if not translated:
        return jsonify({"success": False, "error": "Claude повернув порожній текст"}), 500

    translations[target_lang] = {
        'text': translated,
        'model': result['model'],
        'created_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
    }

    with _get_db() as conn:
        conn.execute(
            'UPDATE transcriptions SET translations_json = ? WHERE id = ?',
            (json.dumps(translations, ensure_ascii=False), transcription_id),
        )
        conn.commit()

    state.metrics.inc("whisper_translate_total", model=result["model"], target_lang=target_lang)
    state.metrics.inc("whisper_translate_tokens_total", value=result["input_tokens"], model=result["model"], kind="input")
    state.metrics.inc("whisper_translate_tokens_total", value=result["output_tokens"], model=result["model"], kind="output")

    return jsonify({
        "success": True,
        "cached": False,
        "translated_text": translated,
        "target_lang": target_lang,
        "model": result['model'],
        "usage": {
            "input_tokens": result['input_tokens'],
            "output_tokens": result['output_tokens'],
            "cache_read_tokens": result['cache_read_tokens'],
        },
    })


@transcription_bp.route('/api/transcription/<int:transcription_id>/topics', methods=['POST'])
def extract_topics_endpoint(transcription_id):
    """Phase 12.19: витягнути topic tags через Claude. Кеш у topics_json."""
    if not text_polishing.is_available():
        return jsonify({"success": False, "error": "ANTHROPIC_API_KEY не встановлено."}), 400

    data = request.get_json(silent=True) or {}
    force = bool(data.get('force', False))
    model_override = data.get('model')

    with _get_db() as conn:
        row = tx_repo.get_by_id(conn, transcription_id, columns=[
            'id', 'transcript_text', 'polished_text', 'segments', 'topics_json',
        ])
        if not row:
            return jsonify({"success": False, "error": "Транскрибування не знайдено"}), 404

        if row['topics_json'] and not force:
            try:
                cached = json.loads(row['topics_json'])
                return jsonify({"success": True, "cached": True, "topics": cached.get('topics', [])})
            except json.JSONDecodeError:
                pass

        segments, has_speakers, speaker_map = _load_speaker_context(conn, transcription_id, row['segments'])

    body_text = row['polished_text'] or row['transcript_text'] or ''
    diarized_text = None
    if has_speakers and speaker_map:
        diarized_text = text_polishing._format_diarized_input(segments, speaker_map)

    if not body_text.strip() and not diarized_text:
        return jsonify({"success": False, "error": "Транскрипт порожній"}), 400

    result, err = _call_claude_endpoint(
        "topics", text_polishing.extract_topics,
        text=body_text, diarized_text=diarized_text, model=model_override,
        error_msg="Помилка Claude API.",
    )
    if err:
        return err

    payload = {"topics": result['topics'], "model": result['model']}
    with _get_db() as conn:
        conn.execute(
            'UPDATE transcriptions SET topics_json = ? WHERE id = ?',
            (json.dumps(payload, ensure_ascii=False), transcription_id),
        )
        conn.commit()

    state.metrics.inc("whisper_topics_total", model=result["model"])

    return jsonify({
        "success": True,
        "cached": False,
        "topics": result['topics'],
        "model": result['model'],
    })


#: Підписи типів коментаря в експортах. Тримаються тут, а не в UI: вигрузку
#: читає людина або інша система, і «correction» англійським слугом посеред
#: українського документа виглядало б як службове сміття.
_CM_LABEL = {
    "correction": "ВИПРАВЛЕННЯ",
    "decision": "РІШЕННЯ",
    "note": "НОТАТКА",
    "context": "КОНТЕКСТ",
    "question": "ПИТАННЯ",
}


def _export_comments(transcription_id) -> list[dict]:
    """Коментарі власника для вигрузки. Порожньо, якщо їх немає, немає id або
    шар ще не мігровано — експорт не має падати через надбудову."""
    if not transcription_id:
        return []
    try:
        from app.services import comments as comments_svc
        return comments_svc.list_for(current_app.config['DATABASE'],
                                     'transcription', int(transcription_id))
    except Exception:
        logger.debug("[export] коментарі не долучено для tx=%s",
                     transcription_id, exc_info=True)
        return []


def _cm_ts(seconds) -> str:
    if seconds is None:
        return ""
    s = int(seconds)
    return f"{s // 60:02d}:{s % 60:02d}"


def _comments_text(rows: list[dict], heading: str = "КОМЕНТАРІ ВЛАСНИКА") -> str:
    """Блок коментарів для текстових вигрузок.

    Іде ПЕРЕД транскриптом, а не додатком у кінці: сенс шару в тому, що
    уточнення власника важить більше за сиру стенограму, і в документі це має
    читатись у тому ж порядку, у якому працює пошук. Явна рамка потрібна, щоб
    у вигрузці не сплуталось, що прозвучало на дзвінку, а що дописано після.
    """
    if not rows:
        return ""
    out = [heading, "=" * len(heading),
           "(написані про запис уже після нього; при суперечності з текстом "
           "нижче правильні саме вони)", ""]
    for c in rows:
        ts = _cm_ts(c.get("anchor_time"))
        head = f"[{_CM_LABEL.get(c.get('kind'), 'НОТАТКА')}]"
        if c.get("pinned"):
            head += " [закріплено]"
        if ts:
            head += f" ~{ts}"
        head += f" · {str(c.get('created_at') or '')[:10]}"
        out.append(head)
        out.append(c.get("body") or "")
        out.append("")
    out.append("-" * 60)
    out.append("")
    return "\n".join(out)


def _comments_md(rows: list[dict]) -> str:
    if not rows:
        return ""
    out = ["## Коментарі власника", "",
           "_Написані про запис уже після нього. При суперечності з "
           "транскриптом нижче правильні саме вони._", ""]
    for c in rows:
        ts = _cm_ts(c.get("anchor_time"))
        bits = [f"**{_CM_LABEL.get(c.get('kind'), 'НОТАТКА')}**"]
        if c.get("pinned"):
            bits.append("закріплено")
        if ts:
            bits.append(f"`{ts}`")
        bits.append(str(c.get("created_at") or "")[:10])
        out.append("- " + " · ".join(bits))
        out.append(f"  {(c.get('body') or '').strip()}")
        out.append("")
    out.append("---")
    out.append("")
    return "\n".join(out)


@transcription_bp.route('/api/export/<format>', methods=['POST'])
def export_transcript(format):
    """Експорт одного транскрипту (TXT/SRT/MD/DOCX/JSON/PDF).

    Phase 10.5: якщо segments мають поле 'speaker' (raw_label) і body містить
    'speakers' (список з resolved name), додаємо префікс 'Ім'я: …' у TXT/SRT.
    Для unnamed labels — fallback 'Спікер N' (по індексу появи).
    """
    data = request.json
    text = data.get('text', '')
    segments = data.get('segments', [])
    # editable-title-description-02: заголовок у контенті та імʼя
    # згенерованого файлу — display_name (title → source_name → «Запис #id»);
    # поле source_name у JSON-payload лишається недоторканим (провенанс).
    # Коли в payload нема ні title, ні source_name, ні id — беремо старий
    # нейтральний фолбек: `display_name` дав би «Запис #None» у заголовку
    # документа й в імені файлу.
    if any(data.get(k) for k in ('title', 'source_name', 'id')):
        record_display_name = record_meta.display_name(data)
    else:
        record_display_name = 'transcript'
    speakers_list = data.get('speakers', []) or []

    # Build raw_label → display_name map (Phase 10.5)
    speaker_display: dict[str, str] = {}
    for s in speakers_list:
        rl = s.get('raw_label')
        if rl and s.get('name'):
            speaker_display[rl] = s['name']
    # Fallback для unnamed: SPEAKER_NN → "Спікер N+1", self → "Ви", UNKNOWN → "?"
    def resolve_speaker(raw_label: str) -> str:
        if not raw_label:
            return ''
        if raw_label in speaker_display:
            return speaker_display[raw_label]
        if raw_label == 'self':
            return 'Ви'
        if raw_label == 'SPEAKER_UNKNOWN':
            return '?'
        m = re.match(r'^SPEAKER_(\d+)$', raw_label)
        if m:
            return f'Спікер {int(m.group(1)) + 1}'
        return raw_label

    has_speakers = any(s.get('speaker') for s in segments)

    # Коментарі власника (Волна 5). Ідуть у txt/md/docx/json — тобто у формати,
    # які людина читає або система переінджестить. У SRT їх немає свідомо:
    # субтитрова доріжка синхронна з аудіо, і врізати туди репліку, якої на
    # плівці не звучало, означало б зіпсувати сам формат.
    export_comments = _export_comments(data.get('id'))

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_source_name = re.sub(r'[^\w\s-]', '', record_display_name)
    safe_source_name = re.sub(r'[-\s]+', '-', safe_source_name)[:50]

    transcripts_dir = os.path.abspath(current_app.config['TRANSCRIPTS_FOLDER'])

    if format == 'txt':
        export_path = os.path.abspath(os.path.join(transcripts_dir, f"{timestamp}_{safe_source_name}.txt"))
        if not export_path.startswith(transcripts_dir):
            logger.error(f"Спроба Path Traversal при експорті: {export_path}")
            return jsonify({"success": False, "error": "Некоректне ім'я файлу"}), 400
        with open(export_path, 'w', encoding='utf-8') as f:
            f.write(_comments_text(export_comments))
            if has_speakers and segments:
                # Group consecutive segments by speaker для читабельності
                prev_speaker = None
                for seg in segments:
                    sp = resolve_speaker(seg.get('speaker', ''))
                    seg_text = seg.get('text', '').strip()
                    if not seg_text:
                        continue
                    if sp != prev_speaker:
                        if prev_speaker is not None:
                            f.write('\n\n')
                        f.write(f'{sp}: ')
                        prev_speaker = sp
                    else:
                        f.write(' ')
                    f.write(seg_text)
                f.write('\n')
            else:
                f.write(text)
    elif format == 'srt':
        export_path = os.path.abspath(os.path.join(transcripts_dir, f"{timestamp}_{safe_source_name}.srt"))
        if not export_path.startswith(transcripts_dir):
            logger.error(f"Спроба Path Traversal при експорті: {export_path}")
            return jsonify({"success": False, "error": "Некоректне ім'я файлу"}), 400
        with open(export_path, 'w', encoding='utf-8') as f:
            for i, segment in enumerate(segments, 1):
                f.write(f"{i}\n")
                f.write(f"{format_srt_timestamp(segment['start'])} --> {format_srt_timestamp(segment['end'])}\n")
                seg_text = segment['text']
                sp_label = segment.get('speaker')
                if has_speakers and sp_label:
                    f.write(f"{resolve_speaker(sp_label)}: {seg_text}\n\n")
                else:
                    f.write(f"{seg_text}\n\n")
    elif format == 'md':
        # Phase 10.9: Markdown — найкращий формат для діалогових транскриптів.
        # Заголовок з summary спікерів, secondary headers per turn з timestamp.
        export_path = os.path.abspath(os.path.join(transcripts_dir, f"{timestamp}_{safe_source_name}.md"))
        if not export_path.startswith(transcripts_dir):
            logger.error(f"Спроба Path Traversal при експорті: {export_path}")
            return jsonify({"success": False, "error": "Некоректне ім'я файлу"}), 400

        def _fmt_ts(seconds: float) -> str:
            mins, secs = divmod(int(seconds), 60)
            hrs, mins = divmod(mins, 60)
            return f'{hrs:02d}:{mins:02d}:{secs:02d}' if hrs else f'{mins:02d}:{secs:02d}'

        with open(export_path, 'w', encoding='utf-8') as f:
            f.write(f'# {record_display_name}\n\n')
            f.write(_comments_md(export_comments))
            if has_speakers and segments:
                # Speaker summary
                from collections import Counter
                spk_counts = Counter(resolve_speaker(s.get('speaker', '')) for s in segments if s.get('text', '').strip())
                spk_line = ', '.join(f'**{name}** ({n})' for name, n in spk_counts.most_common())
                f.write(f'**Спікери:** {spk_line}\n\n---\n\n')

                # Group consecutive segments by speaker
                prev_speaker = None
                turn_text: list[str] = []
                turn_start = 0.0
                def _flush():
                    if prev_speaker and turn_text:
                        f.write(f'### {prev_speaker} · `{_fmt_ts(turn_start)}`\n\n')
                        f.write(' '.join(turn_text))
                        f.write('\n\n')
                for seg in segments:
                    sp = resolve_speaker(seg.get('speaker', ''))
                    seg_text = (seg.get('text') or '').strip()
                    if not seg_text:
                        continue
                    if sp != prev_speaker:
                        _flush()
                        prev_speaker = sp
                        turn_text = [seg_text]
                        turn_start = float(seg.get('start', 0.0))
                    else:
                        turn_text.append(seg_text)
                _flush()
            else:
                # Fallback: plain text без speakers
                f.write(text or '')
                f.write('\n')
    elif format == 'docx':
        # Phase 10.11: DOCX через python-docx — для архіву Word, передачі
        # замовнику, друку. Зберігає speaker structure у вигляді H3-headings.
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
        except ImportError:
            return jsonify({
                "success": False,
                "error": "python-docx не встановлений. .venv/Scripts/python.exe -m pip install python-docx",
            }), 500

        export_path = os.path.abspath(os.path.join(transcripts_dir, f"{timestamp}_{safe_source_name}.docx"))
        if not export_path.startswith(transcripts_dir):
            logger.error(f"Спроба Path Traversal при експорті: {export_path}")
            return jsonify({"success": False, "error": "Некоректне ім'я файлу"}), 400

        def _fmt_ts(seconds: float) -> str:
            mins, secs = divmod(int(seconds), 60)
            hrs, mins = divmod(mins, 60)
            return f'{hrs:02d}:{mins:02d}:{secs:02d}' if hrs else f'{mins:02d}:{secs:02d}'

        doc = Document()
        doc.add_heading(record_display_name or 'Транскрипт', level=1)

        if export_comments:
            doc.add_heading('Коментарі власника', level=2)
            note = doc.add_paragraph()
            note.add_run('Написані про запис уже після нього. При суперечності '
                         'з транскриптом нижче правильні саме вони.').italic = True
            for c in export_comments:
                p = doc.add_paragraph()
                head = _CM_LABEL.get(c.get('kind'), 'НОТАТКА')
                ts = _cm_ts(c.get('anchor_time'))
                p.add_run(f'{head}{" · " + ts if ts else ""}: ').bold = True
                p.add_run(c.get('body') or '')
                p.paragraph_format.space_after = Pt(6)
            doc.add_paragraph()

        if has_speakers and segments:
            from collections import Counter
            spk_counts = Counter(
                resolve_speaker(s.get('speaker', ''))
                for s in segments if (s.get('text') or '').strip()
            )
            summary_p = doc.add_paragraph()
            summary_p.add_run('Спікери: ').bold = True
            spk_parts = [f'{name} ({n})' for name, n in spk_counts.most_common()]
            summary_p.add_run(', '.join(spk_parts))
            doc.add_paragraph()  # spacer

            # Group consecutive same-speaker segments into turns
            prev_speaker = None
            turn_text: list[str] = []
            turn_start = 0.0
            def _flush():
                if prev_speaker and turn_text:
                    h = doc.add_heading(level=3)
                    run = h.add_run(prev_speaker)
                    run.bold = True
                    h.add_run(f'   ·   {_fmt_ts(turn_start)}').italic = True
                    p = doc.add_paragraph(' '.join(turn_text))
                    p.paragraph_format.space_after = Pt(8)
            for seg in segments:
                sp = resolve_speaker(seg.get('speaker', ''))
                seg_text = (seg.get('text') or '').strip()
                if not seg_text:
                    continue
                if sp != prev_speaker:
                    _flush()
                    prev_speaker = sp
                    turn_text = [seg_text]
                    turn_start = float(seg.get('start', 0.0))
                else:
                    turn_text.append(seg_text)
            _flush()
        else:
            doc.add_paragraph(text or '')

        doc.save(export_path)
    elif format == 'json':
        export_path = os.path.abspath(os.path.join(transcripts_dir, f"{timestamp}_{safe_source_name}.json"))
        if not export_path.startswith(transcripts_dir):
            logger.error(f"Спроба Path Traversal при експорті: {export_path}")
            return jsonify({"success": False, "error": "Некоректне ім'я файлу"}), 400
        with open(export_path, 'w', encoding='utf-8') as f:
            # Окремим полем, а не вперемішку з сегментами: споживач JSON має
            # бачити, що це інший шар — свідомо написане про запис, а не
            # розпізнане з нього.
            payload = dict(data)
            payload['display_name'] = record_display_name
            if export_comments:
                payload['comments'] = export_comments
            json.dump(payload, f, ensure_ascii=False, indent=2)
    elif format == 'pdf':
        # Phase 12.13: PDF export via reportlab. Cyrillic-friendly font (Arial з Windows).
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import cm
            from reportlab.lib.colors import HexColor
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
        except ImportError:
            return jsonify({"success": False, "error": "reportlab не встановлений"}), 500

        export_path = os.path.abspath(os.path.join(transcripts_dir, f"{timestamp}_{safe_source_name}.pdf"))
        if not export_path.startswith(transcripts_dir):
            logger.error(f"Спроба Path Traversal при експорті: {export_path}")
            return jsonify({"success": False, "error": "Некоректне ім'я файлу"}), 400

        # Реєструємо Arial для cyrillic. У Windows він є; на інших OS fallback to Helvetica.
        # registerFont — global state; повторні виклики того ж name OK (no-op).
        font_name = 'Helvetica'
        font_bold = 'Helvetica-Bold'
        registered = getattr(pdfmetrics, '_fonts', None) or {}
        for ttf, bold_ttf, name, name_bold in [
            ('C:/Windows/Fonts/arial.ttf', 'C:/Windows/Fonts/arialbd.ttf', 'CyrArial', 'CyrArial-Bold'),
            ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
             '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 'DejaVu', 'DejaVu-Bold'),
        ]:
            if not os.path.exists(ttf):
                continue
            try:
                if name not in registered:
                    pdfmetrics.registerFont(TTFont(name, ttf))
                font_name = name
                if os.path.exists(bold_ttf):
                    if name_bold not in registered:
                        pdfmetrics.registerFont(TTFont(name_bold, bold_ttf))
                    font_bold = name_bold
                else:
                    font_bold = name  # fallback — той же шрифт (без bold variant)
                logger.debug(f"[pdf] fonts: regular={font_name}, bold={font_bold}")
                break
            except Exception as e:
                logger.warning(f"[pdf] font registration failed for {name}: {e}")
                continue

        doc = SimpleDocTemplate(
            export_path, pagesize=A4,
            leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm,
            title=record_display_name, author='Whisper UI',
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'TitleC', parent=styles['Title'],
            fontName=font_bold, fontSize=18, alignment=0, spaceAfter=12,
        )
        meta_style = ParagraphStyle(
            'Meta', parent=styles['Normal'],
            fontName=font_name, fontSize=9, textColor=HexColor('#64748b'), spaceAfter=24,
        )
        speaker_style = ParagraphStyle(
            'Speaker', parent=styles['Normal'],
            fontName=font_bold, fontSize=10, textColor=HexColor('#3B82F6'),
            spaceBefore=10, spaceAfter=2,
        )
        body_style = ParagraphStyle(
            'Body', parent=styles['Normal'],
            fontName=font_name, fontSize=11, leading=15, spaceAfter=6,
        )

        def _esc(s: str) -> str:
            # Reportlab paragraph parses minimal HTML; екрануємо < > & щоб не ламати.
            return (s or '').replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

        story = []
        story.append(Paragraph(_esc(record_display_name), title_style))
        meta_parts = [datetime.now().strftime('%Y-%m-%d %H:%M')]
        if data.get('language'):
            meta_parts.append(f"мова: {data['language']}")
        if data.get('model'):
            meta_parts.append(f"модель: {data['model']}")
        story.append(Paragraph(' · '.join(meta_parts), meta_style))

        if has_speakers and segments:
            prev_speaker = None
            chunk: list[str] = []
            for seg in segments:
                sp = resolve_speaker(seg.get('speaker', ''))
                seg_text = (seg.get('text') or '').strip()
                if not seg_text:
                    continue
                if sp != prev_speaker:
                    if chunk:
                        story.append(Paragraph(_esc(' '.join(chunk)), body_style))
                        chunk = []
                    story.append(Paragraph(_esc(sp), speaker_style))
                    prev_speaker = sp
                chunk.append(seg_text)
            if chunk:
                story.append(Paragraph(_esc(' '.join(chunk)), body_style))
        else:
            for paragraph in (text or '').split('\n\n'):
                if paragraph.strip():
                    story.append(Paragraph(_esc(paragraph), body_style))
                    story.append(Spacer(1, 6))

        doc.build(story)
    else:
        return jsonify({"success": False, "error": "Непідтримуваний формат"}), 400

    response = send_file(export_path, as_attachment=True)

    @response.call_on_close
    def _cleanup_export():
        try:
            if os.path.exists(export_path):
                os.remove(export_path)
        except Exception:
            logger.debug("export cleanup failed for %s (best-effort)", export_path, exc_info=True)
    return response


# ============================================================
# Phase 23 (Story B-B): Video analysis endpoints
# ============================================================

def _allowed_roots_for_media():
    """Return the list of allowed absolute root paths for media serving.

    Covers uploads, youtube_downloads, recordings/sessions (screen video +
    keyframe images written by video_analysis service).
    """
    base = current_app.config.get('BASE_DIR', '.')
    # realpath (not abspath) so symlinks are resolved on both sides of the
    # comparison — prevents a symlinked file inside an allowed root from
    # pointing outside it (defence-in-depth; served paths come from our DB).
    roots = [
        os.path.realpath(current_app.config.get('UPLOAD_FOLDER', '')),
        os.path.realpath(os.path.join(str(base), 'youtube_downloads')),
        os.path.realpath(os.path.join(str(base), 'recordings')),
    ]
    # Also include RECORDING_DIR explicitly (recordings/sessions) in case
    # it is configured to a non-default location.
    rec_dir = current_app.config.get('RECORDING_DIR')
    if rec_dir:
        roots.append(os.path.realpath(str(rec_dir)))
    return [r for r in roots if r]


def _check_path_allowed(path: str) -> bool:
    """Return True if path (resolved via realpath) is under an allowed media root."""
    return safe_path_within_any(_allowed_roots_for_media(), path) is not None


@transcription_bp.route('/api/transcription/<int:tid>/analyze-video', methods=['POST'])
def analyze_video_endpoint(tid):
    """Phase 23 B-B: Submit background video analysis job for a transcription.

    Body (JSON, optional): {"force": bool}
    Returns {"success": false, "reason": "no video"} if no screen recording is
    linked, or {"success": true, "job_id": "<id>"} when the job is queued.
    """
    data = request.get_json(silent=True) or {}
    force = bool(data.get('force', False))
    # Phase 23B: optional per-run vision backend override ('local'|'claude'|'off').
    vision_backend = data.get('vision_backend')
    if vision_backend not in ('local', 'claude', 'off', None):
        vision_backend = None

    try:
        from app.services import video_analysis
    except ImportError as exc:
        logger.error("video_analysis service not available: %s", exc)
        return jsonify({'success': False, 'reason': 'service_unavailable'}), 503

    db_path = current_app.config['DATABASE']

    try:
        video_info = video_analysis.resolve_recording_video(db_path, tid)
    except Exception as exc:
        logger.error("resolve_recording_video failed for tid=%s: %s", tid, exc)
        return jsonify({'success': False, 'reason': 'resolve_error'}), 500

    if video_info is None:
        return jsonify({'success': False, 'reason': 'no video'})

    if state.job_queue is None:
        return jsonify({'success': False, 'reason': 'job_queue_unavailable'}), 503

    def _video_job(job, _db=db_path, _tid=tid, _force=force, _vb=vision_backend):
        return video_analysis.analyze_video(
            _db, _tid, force=_force, job=job, vision_backend=_vb)

    job = state.job_queue.submit(
        'video_analysis',
        _video_job,
        meta={'transcription_id': tid},
    )
    return jsonify({'success': True, 'job_id': job.id})


@transcription_bp.route('/api/transcription/<int:tid>/recording-info', methods=['GET'])
def recording_info(tid):
    """Phase 23 B-B: Return video/recording availability metadata for a transcription.

    Never 500 — degrades gracefully to has_video=False on any error.
    """
    result = {
        'has_video': False,
        'recording_session_id': None,
        'primary_video_available': False,
        'start_offset_sec': 0.0,
        'video_analysis_at': None,
        'keyframes_count': 0,
    }
    try:
        from app.services import video_analysis
        db_path = current_app.config['DATABASE']

        # Fetch video_analysis_at and video_keyframes_count from transcriptions.
        with _get_db() as conn:
            row = conn.execute(
                'SELECT video_analysis_at, video_keyframes_count '
                'FROM transcriptions WHERE id = ?',
                (tid,),
            ).fetchone()
        if row:
            result['video_analysis_at'] = row['video_analysis_at']
            result['keyframes_count'] = row['video_keyframes_count'] or 0

        video_info = video_analysis.resolve_recording_video(db_path, tid)
        if video_info is None:
            return jsonify(result)

        result['has_video'] = True
        result['recording_session_id'] = video_info.get('recording_session_id')
        result['start_offset_sec'] = float(video_info.get('start_offset_sec') or 0.0)

        primary_path = video_info.get('primary_video_path')
        result['primary_video_available'] = bool(
            primary_path and os.path.isfile(primary_path)
        )
    except Exception as exc:
        logger.warning("recording_info error for tid=%s (degrading): %s", tid, exc)

    return jsonify(result)


@transcription_bp.route('/api/transcription/<int:tid>/video', methods=['GET'])
def get_transcription_video(tid):
    """Phase 23 B-B: Stream screen recording video for a transcription.

    Supports HTTP Range requests (seek) via send_file conditional=True.
    """
    try:
        from app.services import video_analysis
        db_path = current_app.config['DATABASE']
        video_info = video_analysis.resolve_recording_video(db_path, tid)
    except Exception as exc:
        logger.error("resolve_recording_video error for tid=%s: %s", tid, exc)
        return jsonify({'success': False, 'error': 'Не вдалося знайти відео'}), 500

    if video_info is None:
        return jsonify({'success': False, 'error': 'Відео не знайдено'}), 404

    primary_path = video_info.get('primary_video_path')
    if not primary_path or not os.path.isfile(primary_path):
        return jsonify({'success': False, 'error': 'Файл відео недоступний'}), 404

    abs_path = os.path.realpath(primary_path)
    if not _check_path_allowed(abs_path):
        logger.warning("Video request blocked (path outside allowed roots): %s", abs_path)
        return jsonify({'success': False, 'error': 'Access denied'}), 403

    return send_file(abs_path, mimetype='video/mp4', conditional=True)


@transcription_bp.route('/api/transcription/<int:tid>/keyframes', methods=['GET'])
def list_keyframes(tid):
    """Phase 23 B-B: List keyframes extracted for a transcription.

    Returns {"keyframes": [{"id", "ts_offset_sec", "ocr_excerpt"}]} ordered
    by ts_offset_sec.  Never 500 — returns empty list on missing table/data.
    """
    try:
        with _get_db() as conn:
            rows = conn.execute(
                'SELECT id, ts_offset_sec, ocr_text, vision_text '
                'FROM video_keyframes '
                'WHERE transcription_id = ? '
                'ORDER BY ts_offset_sec ASC',
                (tid,),
            ).fetchall()
        keyframes = [
            {
                'id': r['id'],
                'ts_offset_sec': r['ts_offset_sec'],
                'ocr_excerpt': (r['ocr_text'] or '')[:80],
                'vision_excerpt': (r['vision_text'] or '')[:200],
            }
            for r in rows
        ]
    except Exception as exc:
        logger.warning("list_keyframes error for tid=%s: %s", tid, exc)
        keyframes = []

    return jsonify({'keyframes': keyframes})


@transcription_bp.route('/api/transcription/<int:tid>/keyframe/<int:kf_id>', methods=['GET'])
def get_keyframe_image(tid, kf_id):
    """Phase 23 B-B: Serve a keyframe image file.

    Validates that the keyframe belongs to this transcription and that the
    image path is within allowed roots before serving.
    """
    try:
        with _get_db() as conn:
            row = conn.execute(
                'SELECT image_path FROM video_keyframes '
                'WHERE id = ? AND transcription_id = ?',
                (kf_id, tid),
            ).fetchone()
    except Exception as exc:
        logger.error("keyframe lookup error kf_id=%s tid=%s: %s", kf_id, tid, exc)
        return jsonify({'success': False, 'error': 'Помилка бази даних'}), 500

    if not row:
        return jsonify({'success': False, 'error': 'Кейфрейм не знайдено'}), 404

    image_path = row['image_path']
    if not image_path or not os.path.isfile(image_path):
        return jsonify({'success': False, 'error': 'Файл зображення недоступний'}), 404

    abs_path = os.path.realpath(image_path)
    if not _check_path_allowed(abs_path):
        logger.warning("Keyframe request blocked (path outside allowed roots): %s", abs_path)
        return jsonify({'success': False, 'error': 'Access denied'}), 403

    return send_file(abs_path, mimetype='image/jpeg', conditional=True)

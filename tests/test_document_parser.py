"""Document parser tests (Phase 16A).

Lightweight, no torch/GPU/live DB - they work on temporary files. Safe even
during active recording/transcription.
"""
import importlib

import pytest

from app.services import document_parser as dp


# ============================================================
# Markdown / TXT
# ============================================================

def test_parse_markdown(tmp_path):
    p = tmp_path / "note.md"
    p.write_text("# Заголовок\n\nПривіт світ. Це тест.", encoding="utf-8")
    res = dp.parse_document(str(p))
    assert res["doc_type"] == "md"
    assert "Заголовок" in res["text"]
    assert "Привіт світ" in res["text"]
    assert res["char_count"] > 0
    assert res["parser_version"] == dp.PARSER_VERSION
    assert len(res["content_hash"]) == 64  # sha256 hex


def test_parse_txt(tmp_path):
    p = tmp_path / "note.txt"
    p.write_text("рядок один\nрядок два", encoding="utf-8")
    res = dp.parse_document(str(p))
    assert res["doc_type"] == "txt"
    assert "рядок один" in res["text"]
    assert res["byte_size"] is not None and res["byte_size"] > 0


def test_normalize_collapses_blank_lines(tmp_path):
    p = tmp_path / "x.txt"
    p.write_text("a\n\n\n\n\nb", encoding="utf-8")
    res = dp.parse_document(str(p))
    assert "\n\n\n" not in res["text"]


def test_filename_override_determines_type(tmp_path):
    # the physical path is .dat, but filename says .md → parsed as text
    p = tmp_path / "blob.dat"
    p.write_text("просто текст", encoding="utf-8")
    res = dp.parse_document(str(p), filename="real.md")
    assert res["doc_type"] == "md"


# ============================================================
# DOCX (python-docx is in the base deps)
# ============================================================

def test_parse_docx(tmp_path):
    docx = pytest.importorskip("docx")
    p = tmp_path / "doc.docx"
    d = docx.Document()
    d.add_heading("Розділ перший", level=1)
    d.add_paragraph("Тіло параграфа з текстом.")
    table = d.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Ключ"
    table.rows[0].cells[1].text = "Значення"
    table.rows[1].cells[0].text = "ціна"
    table.rows[1].cells[1].text = "100"
    d.save(str(p))

    res = dp.parse_document(str(p))
    assert res["doc_type"] == "docx"
    assert "Розділ перший" in res["text"]
    assert "Тіло параграфа" in res["text"]
    # table → markdown
    assert "| Ключ | Значення |" in res["text"]
    assert "| ціна | 100 |" in res["text"]
    assert "---" in res["text"]


# ============================================================
# Errors / edge cases
# ============================================================

def test_unsupported_extension_raises(tmp_path):
    p = tmp_path / "archive.zip"
    p.write_bytes(b"PK\x03\x04")
    with pytest.raises(dp.DocumentParseError):
        dp.parse_document(str(p))


def test_empty_text_raises(tmp_path):
    p = tmp_path / "empty.txt"
    p.write_text("   \n\n  ", encoding="utf-8")
    with pytest.raises(dp.DocumentParseError):
        dp.parse_document(str(p))


def test_is_supported():
    assert dp.is_supported("report.pdf")
    assert dp.is_supported("notes.MD")
    assert dp.is_supported("table.docx")
    assert dp.is_supported("deck.pptx")
    assert dp.is_supported("budget.xlsx")
    assert dp.is_supported("data.CSV")
    assert dp.is_supported("scan.png")
    assert dp.is_supported("photo.JPG")
    assert dp.is_supported("doc.tiff")
    assert not dp.is_supported("song.mp3")
    assert not dp.is_supported("noext")


def test_text_formats_have_no_blocks(tmp_path):
    # md/txt/docx have no structure → blocks=None (chunking falls back to the text)
    p = tmp_path / "n.md"
    p.write_text("# H\n\nтекст", encoding="utf-8")
    res = dp.parse_document(str(p))
    assert res.get("blocks") is None


def test_content_hash_deterministic():
    h1 = dp.content_hash("однаковий текст")
    h2 = dp.content_hash("однаковий текст")
    h3 = dp.content_hash("інший текст")
    assert h1 == h2 != h3


# ============================================================
# PDF - depends on PyMuPDF being present
# ============================================================

def test_pdf_parser_behavior(tmp_path):
    """If PyMuPDF is installed - parse a generated PDF; otherwise ParserUnavailable."""
    has_fitz = importlib.util.find_spec("fitz") is not None
    p = tmp_path / "doc.pdf"
    if not has_fitz:
        p.write_bytes(b"%PDF-1.4 fake")
        with pytest.raises(dp.ParserUnavailable):
            dp.parse_document(str(p))
        return

    # ASCII text: the default base-14 fitz font has no Cyrillic glyphs
    # (insert_text would produce "dots"). Real PDFs embed fonts.
    import fitz
    doc = fitz.open()
    for txt in ("Hello from page one", "Second page content"):
        page = doc.new_page()
        page.insert_text((72, 72), txt)
    doc.save(str(p))
    doc.close()
    res = dp.parse_document(str(p))
    assert res["doc_type"] == "pdf"
    assert res["page_count"] == 2
    assert "Hello" in res["text"]
    # Phase 16B: page blocks with provenance
    blocks = res["blocks"]
    assert blocks and len(blocks) == 2
    assert [b["page"] for b in blocks] == [1, 2]
    assert "page one" in blocks[0]["text"]


def test_parse_pptx_blocks(tmp_path):
    """PPTX → slide blocks with page+section(title) + speaker notes."""
    pptx = pytest.importorskip("pptx")
    from pptx import Presentation
    p = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    slide.shapes.title.text = "Заголовок слайда"
    slide.placeholders[1].text = "Пункт один\nПункт два"
    slide.notes_slide.notes_text_frame.text = "Нотатка доповідача"
    prs.save(str(p))

    res = dp.parse_document(str(p))
    assert res["doc_type"] == "pptx"
    assert res["page_count"] == 1
    blocks = res["blocks"]
    assert blocks and blocks[0]["page"] == 1
    assert blocks[0]["section"] == "Заголовок слайда"
    assert "Пункт один" in res["text"]
    assert "Нотатка доповідача" in res["text"]


def test_pptx_unavailable_raises(tmp_path, monkeypatch):
    """If python-pptx is not installed - ParserUnavailable (graceful)."""
    import importlib
    if importlib.util.find_spec("pptx") is not None:
        pytest.skip("python-pptx is installed - cannot check the unavailable branch")
    p = tmp_path / "deck.pptx"
    p.write_bytes(b"PK\x03\x04 fake pptx")
    with pytest.raises(dp.ParserUnavailable):
        dp.parse_document(str(p))


# ============================================================
# Tables: CSV / XLSX (Phase 16C)
# ============================================================

def test_parse_csv_blocks(tmp_path):
    p = tmp_path / "data.csv"
    p.write_text("Name,Age\nAnn,30\nBob,25\n", encoding="utf-8")
    res = dp.parse_document(str(p))
    assert res["doc_type"] == "csv"
    assert res["page_count"] == 1
    b = res["blocks"][0]
    assert b["page"] == 1 and b["tabular"] is True
    assert b["rows"] == 3 and b["cols"] == 2
    assert "| Name | Age |" in res["text"]
    assert "| Ann | 30 |" in res["text"]


def test_parse_csv_semicolon_delimiter(tmp_path):
    p = tmp_path / "d.csv"
    p.write_text("a;b;c\n1;2;3\n4;5;6\n7;8;9\n", encoding="utf-8")
    res = dp.parse_document(str(p))
    assert "| a | b | c |" in res["text"]


def test_rows_to_markdown_escapes_pipe():
    md, nrows, ncols = dp._rows_to_markdown([["a|b", "c"], ["1", "2"]])
    assert "a\\|b" in md          # pipe escaped
    assert "| --- | --- |" in md  # header separator
    assert nrows == 2 and ncols == 2


def test_md_cell_collapses_whitespace_and_newlines():
    assert dp._md_cell("  multi\nline\ttext ") == "multi line text"
    assert dp._md_cell(None) == ""
    assert dp._md_cell(123) == "123"


def test_parse_xlsx_blocks(tmp_path):
    """XLSX → sheet blocks with section=sheet name + markdown table."""
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook
    p = tmp_path / "book.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Бюджет"
    ws.append(["Категорія", "Сума"])
    ws.append(["Маркетинг", 100])
    ws.append(["Продажі", 50])
    ws2 = wb.create_sheet("Контакти")
    ws2.append(["Імʼя", "Email"])
    ws2.append(["Андрій", "a@x.com"])
    wb.save(str(p))

    res = dp.parse_document(str(p))
    assert res["doc_type"] == "xlsx"
    assert res["page_count"] == 2
    blocks = res["blocks"]
    assert [b["section"] for b in blocks] == ["Бюджет", "Контакти"]
    assert [b["page"] for b in blocks] == [1, 2]
    assert all(b["tabular"] for b in blocks)
    assert "| Категорія | Сума |" in res["text"]
    assert "| Маркетинг | 100 |" in res["text"]
    assert "Аркуш «Бюджет»" in res["text"]


def test_xlsx_unavailable_raises(tmp_path):
    import importlib
    if importlib.util.find_spec("openpyxl") is not None:
        pytest.skip("openpyxl is installed - cannot check the unavailable branch")
    p = tmp_path / "book.xlsx"
    p.write_bytes(b"PK\x03\x04 fake xlsx")
    with pytest.raises(dp.ParserUnavailable):
        dp.parse_document(str(p))


# ============================================================
# OCR: images (Phase 16D)
# ============================================================

def test_ocr_available_returns_bool():
    assert isinstance(dp.ocr_available(), bool)
    assert isinstance(dp.ocr_unavailability_reason(), str)


def test_image_without_ocr_raises(tmp_path):
    if dp.ocr_available():
        pytest.skip("OCR is available - cannot check the unavailable branch")
    p = tmp_path / "scan.png"
    p.write_bytes(b"\x89PNG\r\n\x1a\n fake")
    with pytest.raises(dp.ParserUnavailable):
        dp.parse_document(str(p))


def test_ocr_image_roundtrip(tmp_path):
    """Happy-path OCR: render text into an image and recognize it. Skipped without OCR."""
    if not dp.ocr_available():
        pytest.skip("OCR is unavailable (no pytesseract/tesseract)")
    from PIL import Image, ImageDraw
    p = tmp_path / "scan.png"
    img = Image.new("RGB", (480, 120), "white")
    ImageDraw.Draw(img).text((10, 45), "Hello OCR world", fill="black")
    img.save(str(p))
    res = dp.parse_document(str(p))
    assert res["doc_type"] == "png"
    low = res["text"].lower()
    assert "hello" in low or "ocr" in low
